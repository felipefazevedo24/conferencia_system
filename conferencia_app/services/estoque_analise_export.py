"""Excel da análise dos itens selecionados na tela de Estoque.

Os itens chegam como a tela já os mostrou (mesmo cálculo, sem refazer, e de
qualquer visão); a rota junta o que só o GRV tem: consumo mês a mês e o
histórico de OCs. Consulta que falhou vira aviso na própria planilha."""
from __future__ import annotations

import re
from datetime import date
from io import BytesIO
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import estoque_planejamento_service as planejamento_svc

FMT_QTDE = '#,##0.00'
FMT_CUSTO = '#,##0.0000'
FMT_INT = '#,##0'
FMT_DATA = 'DD/MM/YYYY'
AZUL = "1B3B6F"
SEM_GRV = "Não foi possível consultar o GRV"
CONFIANCA = {"alta": "Alta", "media": "Média", "baixa": "Baixa"}
SITUACAO_OC = {"entregue": "Entregue", "parcial": "Parcial", "aberta": "Em aberto"}

Coluna = tuple[str, str | None, int]


def _num(valor: Any) -> float | None:
    if valor is None or valor == "" or isinstance(valor, bool):
        return None
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def _texto(valor: Any) -> str:
    return str(valor or "").strip()


def _chave(codigo: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(codigo or "").upper())


def _aba(wb: Workbook, titulo: str, colunas: list[Coluna], linhas: list[list[Any]], congelar: str = "C2") -> None:
    ws = wb.create_sheet(titulo)
    ws.append([rotulo for rotulo, _, _ in colunas])
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor=AZUL)
        celula.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for linha in linhas:
        ws.append(linha)
    for indice, (_, formato, largura) in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
        if formato:
            for (celula,) in ws.iter_rows(min_row=2, min_col=indice, max_col=indice):
                celula.number_format = formato
    ws.freeze_panes = congelar
    ws.auto_filter.ref = ws.dimensions


def _colunas_itens(janela: int) -> list[tuple[str, str | None, int, Callable[[dict, dict], Any]]]:
    def saldo(campo):
        # Item que não controla estoque tem saldo zero sempre: vazio, não 0.
        return lambda i, p: None if i.get("sem_controle_estoque") else _num(i.get(campo))

    def consumo_mensal(i, p):
        if p:
            return _num(p.get("consumo_mensal"))
        diario = _num(i.get("consumo_diario"))
        return diario * planejamento_svc.DIAS_POR_MES if diario else None

    def ocs(i, p):
        return "; ".join(
            f"OC {_texto(o.get('ordem_compra'))} · {_texto(o.get('fornecedor'))} · {_num(o.get('quantidade_pendente')) or 0:g}"
            for o in i.get("ordens_compra_abertas") or []
        )

    def lead_fonte(i, p):
        if not p:
            return None
        return "Histórico" if p.get("lead_time_fonte") == "historico" else "Padrão"

    return [
        ("Código", None, 14, lambda i, p: _texto(i.get("codigo"))),
        ("Descrição", None, 46, lambda i, p: _texto(i.get("descricao"))),
        ("Visão", None, 18, lambda i, p: _texto(i.get("visao"))),
        ("Un", None, 6, lambda i, p: _texto(i.get("unidade"))),
        ("Família", None, 28, lambda i, p: _texto(i.get("familia"))),
        ("Status", None, 24, lambda i, p: _texto(i.get("nivel_label"))),
        ("Disponível", FMT_QTDE, 12, saldo("saldo_disponivel")),
        ("Reservado", FMT_QTDE, 12, saldo("saldo_reservado")),
        ("Total no depósito", FMT_QTDE, 12, saldo("saldo_total")),
        ("Localização", None, 18, lambda i, p: ", ".join(_texto(x) for x in i.get("localizacoes") or [])),
        ("Consumo/mês", FMT_QTDE, 12, consumo_mensal),
        (f"Saída em {janela} dias", FMT_QTDE, 12, lambda i, p: _num(i.get("saida_total_periodo"))),
        (f"Dias com saída ({janela} dias)", FMT_INT, 12, lambda i, p: _num(i.get("dias_com_saida"))),
        (f"Maior saída ({janela} dias)", FMT_QTDE, 12, lambda i, p: _num(i.get("maior_saida"))),
        ("Data da maior saída", FMT_DATA, 12, lambda i, p: planejamento_svc._data(i.get("maior_saida_data"))),
        ("Cobertura (dias)", FMT_INT, 11, lambda i, p: _num(i.get("cobertura_dias"))),
        ("Cobertura com OC (dias)", FMT_INT, 12, lambda i, p: _num(i.get("cobertura_projetada_dias"))),
        ("Cobre retirada típica (%)", FMT_QTDE, 12, lambda i, p: _num(p.get("cobre_retirada_pct"))),
        ("Pendente em OC", FMT_QTDE, 12, lambda i, p: _num(i.get("quantidade_pendente_oc")) or None),
        ("OCs abertas", None, 40, ocs),
        ("Curva ABC", None, 8, lambda i, p: _texto(p.get("classe_abc"))),
        ("Nível de serviço (%)", FMT_QTDE, 11, lambda i, p: _num(p.get("nivel_servico_pct"))),
        ("Tipo de demanda", None, 14, lambda i, p: _texto(p.get("tipo_demanda_label"))),
        ("Meses com saída (24 meses)", FMT_INT, 12, lambda i, p: _num(p.get("meses_com_saida_24m"))),
        ("Meses analisados", FMT_INT, 11, lambda i, p: _num(p.get("meses_analisados_24m"))),
        ("Meses com saída desde o início", FMT_INT, 13, lambda i, p: _num(p.get("meses_com_saida_total"))),
        ("Primeira movimentação", FMT_DATA, 13, lambda i, p: planejamento_svc._data(p.get("primeira_movimentacao"))),
        ("Consumo 12 meses", FMT_QTDE, 12, lambda i, p: _num(p.get("consumo_12m"))),
        ("Desvio mensal", FMT_QTDE, 11, lambda i, p: _num(p.get("desvio_mensal"))),
        ("Intervalo médio (ADI)", FMT_QTDE, 11, lambda i, p: _num(p.get("adi"))),
        ("Variação (CV²)", FMT_QTDE, 11, lambda i, p: _num(p.get("cv2"))),
        ("Lead time (dias)", FMT_QTDE, 11, lambda i, p: _num(p.get("lead_time_dias"))),
        ("Desvio do lead time", FMT_QTDE, 11, lambda i, p: _num(p.get("lead_time_desvio"))),
        ("Entregas no lead time", FMT_INT, 11, lambda i, p: _num(p.get("lead_time_amostras"))),
        ("Origem do lead time", None, 12, lead_fonte),
        ("Retirada típica", FMT_QTDE, 12, lambda i, p: _num(p.get("retirada_tipica"))),
        ("Retiradas (24 meses)", FMT_INT, 11, lambda i, p: _num(p.get("retiradas_24m"))),
        ("Estoque de segurança", FMT_QTDE, 12, lambda i, p: _num(p.get("estoque_seguranca"))),
        ("Mínimo sugerido", FMT_QTDE, 12, lambda i, p: _num(p.get("estoque_minimo_sugerido"))),
        ("Mínimo no GRV", FMT_QTDE, 12, lambda i, p: _num(i.get("estoque_minimo")) or None),
        ("Lote sugerido", FMT_QTDE, 12, lambda i, p: _num(p.get("lote_economico_sugerido"))),
        ("Lote no GRV", FMT_QTDE, 12, lambda i, p: _num(i.get("lote_economico")) or None),
        ("Custo unitário", FMT_CUSTO, 12, lambda i, p: _num(p.get("custo_unitario"))),
        ("Confiança", None, 10, lambda i, p: CONFIANCA.get(p.get("confianca"), "")),
        ("Motivos da confiança", None, 34, lambda i, p: " · ".join(_texto(m) for m in p.get("confianca_motivos") or [])),
        ("Consumo para OS em 12 meses (%)", FMT_QTDE, 13, lambda i, p: _num(p.get("consumo_os_pct_12m"))),
        ("OCs para OS (12 meses)", FMT_INT, 11, lambda i, p: _num(p.get("oc_para_os_12m"))),
        ("OCs para estoque (12 meses)", FMT_INT, 11, lambda i, p: _num(p.get("oc_estoque_12m"))),
        ("Candidato a item de estoque", None, 12, lambda i, p: ("Sim" if p.get("candidato_estoque") else "Não") if p else ""),
        ("Reservas", FMT_INT, 9, lambda i, p: len(i.get("reservas") or []) or None),
    ]


def gerar_excel_analise(
    itens: list[dict[str, Any]],
    series: dict[str, Any] | None,
    fornecimentos: dict[str, dict[str, Any] | None],
    *,
    janela_dias: int,
    hoje: date,
) -> BytesIO:
    """`series`: histórico do planejamento por código sem pontuação (None = GRV
    fora). `fornecimentos`: resumir_fornecimentos() por código (None = falhou)."""
    wb = Workbook()
    wb.remove(wb.active)

    colunas = _colunas_itens(janela_dias)
    _aba(
        wb, "Itens",
        [(rotulo, formato, largura) for rotulo, formato, largura, _ in colunas],
        [[obter(item, item.get("planejamento") or {}) for _, _, _, obter in colunas] for item in itens],
    )

    # Consumo mês a mês: os 24 meses fechados que o cálculo usa + o mês corrente.
    meses = planejamento_svc._meses_fechados(hoje, planejamento_svc.MESES_CLASSIFICACAO)
    mes_atual = f"{hoje.year:04d}-{hoje.month:02d}"
    colunas_consumo: list[Coluna] = [("Código", None, 14), ("Descrição", None, 46), ("Un", None, 6)]
    colunas_consumo += [(f"{m[5:]}/{m[:4]}", FMT_QTDE, 10) for m in meses]
    colunas_consumo += [(f"{mes_atual[5:]}/{mes_atual[:4]} (parcial)", FMT_QTDE, 11), ("Total 24 meses", FMT_QTDE, 12)]
    linhas_consumo = []
    for item in itens:
        base = [_texto(item.get("codigo")), _texto(item.get("descricao")), _texto(item.get("unidade"))]
        historico = (series or {}).get(_chave(item.get("codigo")))
        if not historico:
            linhas_consumo.append(base + [SEM_GRV if series is None else "Sem histórico no GRV"])
            continue
        saidas = {m.get("mes"): _num(m.get("saida")) or 0.0 for m in historico.get("serie_mensal") or []}
        inicio = str(historico.get("primeira_movimentacao") or "")[:7]
        # Antes do primeiro movimento o item nem existia: vazio, não zero.
        valores = [None if inicio and m < inicio else saidas.get(m, 0.0) for m in meses + [mes_atual]]
        linhas_consumo.append(base + valores + [sum(v or 0 for v in valores[:-1])])
    _aba(wb, "Consumo mensal", colunas_consumo, linhas_consumo, congelar="D2")

    linhas_forn, linhas_compras = [], []
    meses_atual = planejamento_svc.MESES_FORNECEDOR_ATUAL
    for item in itens:
        codigo = _texto(item.get("codigo"))
        base = [codigo, _texto(item.get("descricao"))]
        dados = fornecimentos.get(codigo)
        if dados is None:
            linhas_forn.append(base + [SEM_GRV])
            linhas_compras.append(base + [None, None, SEM_GRV])
            continue
        if not dados.get("fornecimentos"):
            linhas_forn.append(base + ["Nenhuma OC desde 2019"])
            continue
        for f in dados.get("fornecedores") or []:
            selo = "Mais rápido e menor preço" if f.get("recomendado") else ("Mais rápido" if f.get("mais_rapido") else ("Menor preço" if f.get("menor_preco") else ""))
            linhas_forn.append(base + [
                _texto(f.get("fornecedor")), _num(f.get("fornecimentos")), _num(f.get("lead_time_medio")),
                _num(f.get("entregas_com_prazo")), _num(f.get("ultimo_preco")), planejamento_svc._data(f.get("ultimo_preco_data")),
                planejamento_svc._data(f.get("ultima_compra")), "Sim" if f.get("atual") else "Não", selo,
            ])
        unidade = _texto(item.get("unidade"))
        for f in dados["fornecimentos"]:
            linhas_compras.append(base + [
                planejamento_svc._data(f.get("dt_oc")), _texto(f.get("ordem_compra")), _texto(f.get("fornecedor")),
                SITUACAO_OC.get(f.get("situacao"), ""), planejamento_svc._data(f.get("dt_entrada")),
                _num(f.get("quantidade")), unidade, _num(f.get("preco_unidade_estoque")),
                _num(f.get("quantidade_compra")), _texto(f.get("unidade_compra")), _num(f.get("preco_compra")),
                _num(f.get("lead_time_dias")), "Sim" if f.get("para_os") else "Não",
            ])
    _aba(wb, "Fornecedores", [
        ("Código", None, 14), ("Descrição", None, 46), ("Fornecedor", None, 40), ("Compras", FMT_INT, 9),
        ("Prazo médio (dias)", FMT_QTDE, 11), ("Entregas com prazo", FMT_INT, 11), ("Último preço/un", FMT_CUSTO, 13),
        ("Data do último preço", FMT_DATA, 12), ("Última compra", FMT_DATA, 12),
        (f"Comprou nos últimos {meses_atual} meses", None, 13), ("Destaque", None, 24),
    ], linhas_forn)
    _aba(wb, "Compras", [
        ("Código", None, 14), ("Descrição", None, 46), ("Data da OC", FMT_DATA, 12), ("OC", None, 9),
        ("Fornecedor", None, 40), ("Situação", None, 11), ("Entrada", FMT_DATA, 12), ("Quantidade", FMT_QTDE, 12),
        ("Un", None, 6), ("Preço/un", FMT_CUSTO, 12), ("Qtde na un. de compra", FMT_QTDE, 12),
        ("Un. de compra", None, 9), ("Preço na un. de compra", FMT_CUSTO, 13), ("Prazo (dias)", FMT_INT, 9),
        ("Para OS", None, 8),
    ], linhas_compras)

    linhas_reservas = []
    for item in itens:
        for r in item.get("reservas") or []:
            linhas_reservas.append([
                _texto(item.get("codigo")), _texto(item.get("descricao")), _texto(item.get("unidade")),
                _num(r.get("quantidade")), _texto(r.get("os")), _texto(r.get("cod_os")), _texto(r.get("titulo_os")),
                _texto(r.get("orcamento")), _texto(r.get("versao")), _texto(r.get("cliente")),
            ])
    _aba(wb, "Reservas", [
        ("Código", None, 14), ("Descrição", None, 46), ("Un", None, 6), ("Quantidade", FMT_QTDE, 12),
        ("OS", None, 12), ("Cód. OS", None, 10), ("Título da OS", None, 36), ("Orçamento", None, 12),
        ("Versão", None, 8), ("Cliente", None, 36),
    ], linhas_reservas)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
