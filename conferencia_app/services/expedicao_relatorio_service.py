"""Relatório da Conferência de Expedição em Excel (abas FAT e ST juntas).

Cada número usa a data do próprio evento: conferida pela data da conferência,
faturada pelo faturado_at, expedida pelo expedido_at. Uma ordem pode entrar em
mais de um número.

A data e o conferente da conferência saem da trilha (ExpedicaoConferenciaLog):
editar a conferência regrava conferido_at e conferente na ordem, então a ordem
sozinha mostra a última edição, não quem conferiu nem quando. Ordem antiga, sem
trilha, cai nos campos da ordem.

Divergência: conferência com contagem errada não é salva (o sistema bloqueia e
a pessoa reconta), então ordem.divergente fica sempre falso. O que sobra é a
tentativa bloqueada (trilha) e o conferido x NF (só FAT).

Conferência cega: a quantidade esperada nunca sai daqui.
"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, or_

from ..extensions import db
from ..models import (
    ExpedicaoConferenciaLog,
    ExpedicaoFatComparacaoNF,
    ExpedicaoOrdemFat,
    ExpedicaoOrdemFatItem,
    ExpedicaoOrdemST,
    ExpedicaoOrdemSTItem,
)
from ..tempo import agora_br
from . import expedicao_fat_comparacao_nf_service as comparacao_nf_svc
from . import expedicao_fat_service as fat_svc
from .viagem_relatorio_service import _FMT_DATA, _SEM_INFO, _append, _data, _lista, _secao, _tabela

TIPOS = {"fat": "FAT", "st": "ST"}
STATUS = {
    fat_svc.status_slug(s): s
    for s in (
        fat_svc.STATUS_PENDENTE, fat_svc.STATUS_CONFERIDO, fat_svc.STATUS_FATURADO_SEM_CONF, fat_svc.STATUS_FATURADO,
        fat_svc.STATUS_EM_ROMANEIO, fat_svc.STATUS_EXPEDIDO, fat_svc.STATUS_FINALIZADO_SEM_CONF,
    )
}
_ACOES = (
    "conferencia", "edicao", "seguir_sem_contagem", "estorno_conferencia",
    "tentativa_bloqueada", "finalizacao_sem_conferencia", "expedicao_sem_conferencia",
)
_SEM_CONTAGEM = "Sem contagem (Admin)"
_FMT_DIAS = "#,##0.0"


def opcoes() -> dict:
    """Listas do painel: quem já conferiu e os clientes/fornecedores."""
    conferentes, parceiros = set(), set()
    for modelo, parceiro in ((ExpedicaoOrdemFat, ExpedicaoOrdemFat.cliente), (ExpedicaoOrdemST, ExpedicaoOrdemST.fornecedor)):
        for conferente, nome in db.session.query(modelo.conferente, parceiro).filter(modelo.excluido.is_(False)).distinct():
            conferentes.add((conferente or "").strip())
            parceiros.add((nome or "").strip())
    usuarios = db.session.query(ExpedicaoConferenciaLog.usuario).filter(
        ExpedicaoConferenciaLog.acao.in_(("conferencia", "seguir_sem_contagem", "tentativa_bloqueada"))
    ).distinct()
    conferentes.update((u or "").strip() for (u,) in usuarios)
    conferentes.discard("")
    parceiros.discard("")
    return {
        "conferentes": sorted(conferentes, key=str.casefold),
        "parceiros": sorted(parceiros, key=str.casefold),
        "status": [{"valor": k, "nome": v} for k, v in STATUS.items()],
    }


def montar_filtros(args) -> dict:
    inicio = _data(args.get("inicio"), "Informe a data inicial.")
    fim = _data(args.get("fim"), "Informe a data final.")
    if fim < inicio:
        raise ValueError("A data final é anterior à inicial.")
    tipo = str(args.get("tipo") or "").strip().lower()
    if tipo and tipo not in TIPOS:
        raise ValueError("Tipo inválido.")
    status = str(args.get("status") or "").strip()
    if status and status not in STATUS:
        raise ValueError("Status inválido.")
    return {
        "inicio": inicio,
        "fim": fim,
        "tipo": tipo,
        "conferente": str(args.get("conferente") or "").strip()[:100],
        "parceiro": str(args.get("parceiro") or "").strip()[:160],
        "status": status,
        "so_divergentes": str(args.get("divergentes") or "").strip().lower() in ("1", "true", "on", "sim"),
    }


def gerar_excel(args, usuario: str) -> tuple[BytesIO, str]:
    filtros = montar_filtros(args)
    dados = coletar(filtros)
    wb = Workbook()
    _aba_resumo(wb.active, filtros, dados, usuario)
    _aba_ordens(wb.create_sheet("Ordens"), dados["ordens"])
    _aba_divergencias(wb.create_sheet("Divergências"), dados["divergencias"])
    _aba_sem_conferencia(wb.create_sheet("Sem conferência"), dados["sem_conferencia"])
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    nome = f"relatorio_conferencia_expedicao_{filtros['inicio']:%Y%m%d}_{filtros['fim']:%Y%m%d}.xlsx"
    return stream, nome


# --------------------------------------------------------------------------- dados

def coletar(f: dict) -> dict:
    ini = datetime.combine(f["inicio"], time.min)
    fim = datetime.combine(f["fim"] + timedelta(days=1), time.min)

    def no_periodo(valor) -> bool:
        return valor is not None and ini <= valor < fim

    ordens = []
    for origem in TIPOS:
        if not f["tipo"] or f["tipo"] == origem:
            ordens += _ordens(origem, f, ini, fim)

    resultado = []
    divergencias, sem_conferencia = [], []
    for o in ordens:
        o["tentativas_periodo"] = [t for t in o["tentativas"] if no_periodo(t["data"])]
        o["finalizacoes_periodo"] = [t for t in o["finalizacoes"] if no_periodo(t["data"])]
        o["conferida_no_periodo"] = no_periodo(o["conferida_em"])
        o["faturada_no_periodo"] = no_periodo(o["faturada_em"])
        o["expedida_no_periodo"] = no_periodo(o["expedida_em"])
        if not (o["conferida_no_periodo"] or o["faturada_no_periodo"] or o["expedida_no_periodo"]
                or o["tentativas_periodo"] or o["finalizacoes_periodo"]):
            # Entrou só por edição/alteração no período, sem evento que conte.
            continue
        if f["conferente"] and o["conferente"] != f["conferente"]:
            continue
        if f["so_divergentes"] and not o["divergente"]:
            continue
        resultado.append(o)

        base = {"tipo": o["tipo"], "ordem": o["ordem"], "parceiro": o["parceiro"]}
        for t in o["tentativas_periodo"]:
            for item in t["itens"] or [{}]:
                divergencias.append({**base, "data": t["data"], "divergencia": "Contagem bloqueada", "quem": t["usuario"],
                                     "codigo": item.get("cod_interno") or "", "item": item.get("item") or "",
                                     "os": item.get("n_os") or "", "qtd": item.get("para"), "qtd_nf": None})
        if o["faturada_no_periodo"]:
            for d in o["divergencias_nf"]:
                divergencias.append({**base, "data": o["faturada_em"], "divergencia": f"NF: {d.get('tipo_label') or ''}",
                                     "quem": o["conferente"], "codigo": d.get("cod_interno") or "", "item": d.get("item") or "",
                                     "os": "", "qtd": d.get("conferido"), "qtd_nf": d.get("faturado")})

        base["nf"] = o["nf"]
        if o["conferida_no_periodo"] and o["como"] == _SEM_CONTAGEM:
            sem_conferencia.append({**base, "data": o["conferida_em"], "situacao": "Seguiu sem contagem (Admin)",
                                    "quem": o["conferente"], "motivo": o["motivo_sem_contagem"]})
        for t in o["finalizacoes_periodo"]:
            sem_conferencia.append({**base, "data": t["data"], "situacao": "Finalizada sem conferência (Admin)",
                                    "quem": t["usuario"], "motivo": t["motivo"]})
        if o["expedida_no_periodo"] and o["expedida_sem_conf"]:
            sem_conferencia.append({**base, "data": o["expedida_em"], "situacao": "Expedida sem conferência",
                                    "quem": o["liberada_sem_conf_por"], "motivo": o["motivo_expedida"]})

    resultado.sort(key=lambda o: (o["tipo"], min(filter(None, [
        o["conferida_em"] if o["conferida_no_periodo"] else None,
        o["faturada_em"] if o["faturada_no_periodo"] else None,
        o["expedida_em"] if o["expedida_no_periodo"] else None,
        *(t["data"] for t in o["tentativas_periodo"] + o["finalizacoes_periodo"]),
    ]))))
    divergencias.sort(key=lambda d: d["data"] or datetime.min)
    sem_conferencia.sort(key=lambda d: d["data"] or datetime.min)
    return {"ordens": resultado, "divergencias": divergencias, "sem_conferencia": sem_conferencia}


def _ordens(origem: str, f: dict, ini: datetime, fim: datetime) -> list[dict]:
    fat = origem == "fat"
    modelo, item_modelo = (ExpedicaoOrdemFat, ExpedicaoOrdemFatItem) if fat else (ExpedicaoOrdemST, ExpedicaoOrdemSTItem)

    def periodo(col):
        return and_(col >= ini, col < fim)

    # Conferência, tentativa e finalização têm a data na trilha: a ordem com
    # um desses no período entra mesmo que o campo dela tenha outra data.
    ids_trilha = [oid for (oid,) in db.session.query(ExpedicaoConferenciaLog.ordem_id).filter(
        ExpedicaoConferenciaLog.origem == origem,
        ExpedicaoConferenciaLog.acao.in_(("conferencia", "seguir_sem_contagem", "tentativa_bloqueada", "finalizacao_sem_conferencia")),
        periodo(ExpedicaoConferenciaLog.created_at),
    ).distinct()]
    condicoes = [periodo(modelo.conferido_at), periodo(modelo.faturado_at), periodo(modelo.expedido_at)]
    if ids_trilha:
        condicoes.append(modelo.id.in_(ids_trilha))
    registros = modelo.query.filter(modelo.excluido.is_(False), or_(*condicoes)).all()
    if f["parceiro"]:
        # Em Python: o lower() do SQLite não mexe em acento.
        busca = f["parceiro"].casefold()
        registros = [o for o in registros if busca in ((o.cliente if fat else o.fornecedor) or "").casefold()]
    if f["status"]:
        registros = [o for o in registros if fat_svc.status_slug(o.status) == f["status"]]
    if fat:
        # Mesmo corte da lista da tela: o backlog antigo faturado sem
        # conferência não aparece lá, então não conta aqui.
        from ..routes.expedicao_fat_routes import FAT_SEM_CONF_COD_MINIMO

        registros = [
            o for o in registros
            if not (fat_svc.status_slug(o.status) == "faturado_sem_conf" and (o.cod_ordem_fat or 0) < FAT_SEM_CONF_COD_MINIMO)
        ]
    if not registros:
        return []
    ids = [o.id for o in registros]

    itens = {
        oid: (linhas, contadas, pecas)
        for oid, linhas, contadas, pecas in db.session.query(
            item_modelo.ordem_id, func.count(item_modelo.id), func.count(item_modelo.qtde_conferida), func.sum(item_modelo.qtde_conferida),
        ).filter(item_modelo.ordem_id.in_(ids)).group_by(item_modelo.ordem_id)
    }
    trilha = defaultdict(list)
    for lg in ExpedicaoConferenciaLog.query.filter(
        ExpedicaoConferenciaLog.origem == origem,
        ExpedicaoConferenciaLog.ordem_id.in_(ids),
        ExpedicaoConferenciaLog.acao.in_(_ACOES),
    ).order_by(ExpedicaoConferenciaLog.created_at, ExpedicaoConferenciaLog.id):
        trilha[lg.ordem_id].append(lg)

    divergentes_nf, detalhes_nf = set(), {}
    if fat:
        divergentes_nf = comparacao_nf_svc.ordens_divergentes(registros)
        if divergentes_nf:
            for linha in ExpedicaoFatComparacaoNF.query.filter(ExpedicaoFatComparacaoNF.ordem_id.in_(divergentes_nf)):
                detalhes_nf[linha.ordem_id] = _json(linha.detalhes) or []

    linhas = []
    for o in registros:
        linhas_item, contadas, pecas = itens.get(o.id, (0, 0, 0))
        conf = _da_trilha(trilha[o.id], o, contadas)
        faturada_em = o.faturado_at
        if not faturada_em:
            conf_x_fat = ""
        elif not conf["conferida_em"]:
            conf_x_fat = "Sem conferência"
        else:
            conf_x_fat = "Depois do faturamento" if conf["conferida_em"] > faturada_em else "Antes do faturamento"
        linhas.append({
            "tipo": TIPOS[origem],
            "codigo_interno": o.codigo_interno or "",
            "ordem": str(o.cod_ordem_fat) if fat else (o.cod_ordem_compra or ""),
            "parceiro": ((o.cliente if fat else o.fornecedor) or "").strip(),
            "orcamento": (o.orcamento or "") if fat else "",
            "pedido": (o.pedido or "") if fat else "",
            "os": "" if fat else (o.n_os or ""),
            "nf": o.numero_nf or "",
            "status": o.status or "",
            "operacao": (o.operacao_tipo or "nacional").capitalize() if fat else "",
            "solicitada_em": o.dt_solicitacao_fat if fat else o.dt_solicitacao,
            "previsao": o.dt_previsao_entrega if fat else o.dt_prevista_entrega,
            **conf,
            "conf_x_fat": conf_x_fat,
            "faturada_em": faturada_em,
            "expedida_em": o.expedido_at,
            "expedida_por": o.expedido_by or "",
            "expedida_sem_conf": bool(o.expedido_sem_conferencia),
            "motivo_expedida": o.expedido_sem_conferencia_motivo or "",
            "nf_divergente": o.id in divergentes_nf,
            "divergencias_nf": detalhes_nf.get(o.id, []),
            "divergente": bool(conf["tentativas"]) or o.id in divergentes_nf,
            "linhas_item": linhas_item,
            # Contagem só depois de conferida: antes disso é progresso salvo.
            "pecas": int(pecas or 0) if conf["conferida_em"] and conf["como"] != _SEM_CONTAGEM else None,
            "volumes": o.qtde_volumes or "",
            "especie": o.especie_volumes or "",
            "peso_bruto": o.peso_bruto or "",
            "peso_liquido": o.peso_liquido or "",
        })
    return linhas


def _da_trilha(logs, ordem, contadas: int) -> dict:
    """Data, conferente e jeito da conferência atual, mais tentativas e
    exceções. Estorno zera: a conferência que vale é a primeira depois dele."""
    conferida_em = conferente = None
    como = motivo_sem_contagem = ""
    edicoes = 0
    tentativas, finalizacoes = [], []
    liberada_por = ""
    for lg in logs:
        detalhes = _json(lg.detalhes) or {}
        if lg.acao in ("conferencia", "seguir_sem_contagem"):
            if conferida_em is None:
                conferida_em, conferente = lg.created_at, lg.usuario
            como = _SEM_CONTAGEM if lg.acao == "seguir_sem_contagem" else "Contagem"
            motivo_sem_contagem = _motivo(detalhes) if lg.acao == "seguir_sem_contagem" else ""
        elif lg.acao == "edicao":
            edicoes += 1
            como = "Contagem"
        elif lg.acao == "estorno_conferencia":
            conferida_em = conferente = None
            como = motivo_sem_contagem = ""
            edicoes = 0
        elif lg.acao == "tentativa_bloqueada":
            itens = [i for i in (detalhes.get("itens") or []) if i.get("divergente")]
            tentativas.append({"data": lg.created_at, "usuario": lg.usuario or "", "itens": itens})
        elif lg.acao == "finalizacao_sem_conferencia":
            finalizacoes.append({"data": lg.created_at, "usuario": lg.usuario or "", "motivo": _motivo(detalhes)})
        elif lg.acao == "expedicao_sem_conferencia":
            liberada_por = lg.usuario or ""

    if not ordem.conferido_at:
        conferida_em = conferente = None
        como = motivo_sem_contagem = ""
    elif conferida_em is None:
        # Conferida antes da trilha existir: fica o que a ordem guardou.
        conferida_em, conferente = ordem.conferido_at, ordem.conferente
        como = "Contagem" if contadas else _SEM_CONTAGEM
    return {
        "conferida_em": conferida_em,
        "conferente": (conferente or "").strip(),
        "como": como,
        "motivo_sem_contagem": motivo_sem_contagem,
        "edicoes": edicoes,
        "tentativas": tentativas,
        "finalizacoes": finalizacoes,
        "liberada_sem_conf_por": liberada_por,
    }


def _motivo(detalhes: dict) -> str:
    for mudanca in detalhes.get("cabecalho") or []:
        if mudanca.get("campo") == "motivo":
            return str(mudanca.get("para") or "")
    return ""


def _json(texto):
    try:
        return json.loads(texto) if texto else None
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------- planilha

def _aba_resumo(ws, f, dados, usuario):
    ws.title = "Resumo"
    ordens = dados["ordens"]
    conferidas = [o for o in ordens if o["conferida_no_periodo"]]
    faturadas = [o for o in ordens if o["faturada_no_periodo"]]
    expedidas = [o for o in ordens if o["expedida_no_periodo"]]
    sem_contagem = [o for o in conferidas if o["como"] == _SEM_CONTAGEM]
    com_contagem = [o for o in conferidas if o["como"] != _SEM_CONTAGEM]
    tentativas = [t for o in ordens for t in o["tentativas_periodo"]]

    ws.append(["Relatório da Conferência de Expedição"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Período: {f['inicio']:%d/%m/%Y} a {f['fim']:%d/%m/%Y}"])
    ws.append([f"Filtros: {_descrever_filtros(f)}"])
    ws.append([f"Gerado em {agora_br():%d/%m/%Y %H:%M} por {usuario}"])
    ws.append([])

    _secao(ws, "Números do período", ["", "Total"])
    numeros = [
        ("Ordens conferidas", len(conferidas), None),
        ("   com contagem física", len(com_contagem), None),
        ("   sem contagem (liberadas por Admin)", len(sem_contagem), None),
        ("   conferidas depois de faturadas", sum(1 for o in conferidas if o["conf_x_fat"] == "Depois do faturamento"), None),
        ("Linhas de item conferidas", sum(o["linhas_item"] for o in com_contagem), None),
        ("Peças contadas", sum(o["pecas"] or 0 for o in com_contagem), None),
        ("Tentativas bloqueadas (contagem não bateu)", len(tentativas), None),
        ("Ordens com tentativa bloqueada", sum(1 for o in ordens if o["tentativas_periodo"]), None),
        ("Finalizadas sem conferência (Admin)", sum(len(o["finalizacoes_periodo"]) for o in ordens), None),
        ("Ordens faturadas", len(faturadas), None),
        ("   faturadas antes da conferência", sum(1 for o in faturadas if o["conf_x_fat"] != "Antes do faturamento"), None),
        ("   conferido diferente da NF (só FAT)", sum(1 for o in faturadas if o["nf_divergente"]), None),
        ("Ordens expedidas", len(expedidas), None),
        ("   expedidas sem conferência", sum(1 for o in expedidas if o["expedida_sem_conf"]), None),
        ("Dias da solicitação até a conferência (média)",
         _media_dias((o["solicitada_em"], o["conferida_em"]) for o in conferidas), _FMT_DIAS),
        ("Dias da conferência até a expedição (média)",
         _media_dias((o["conferida_em"], o["expedida_em"]) for o in expedidas), _FMT_DIAS),
    ]
    for rotulo, valor, formato in numeros:
        _append(ws, [rotulo, valor if valor is not None else "Sem dado"])
        if formato and valor is not None:
            ws.cell(row=ws.max_row, column=2).number_format = formato
    ws.append([])

    conf_cab = ["", "Conferidas", "Sem contagem", "Linhas de item", "Peças contadas", "Tentativas bloqueadas"]
    _tabela(ws, "Por conferente", conf_cab, _quebra_conferente(ordens))
    fluxo_cab = ["", "Conferidas", "Faturadas", "Expedidas", "Com divergência"]
    _tabela(ws, "Por tipo", fluxo_cab, _quebra_fluxo(ordens, lambda o: o["tipo"]))
    _tabela(ws, "Por cliente / fornecedor", fluxo_cab, _quebra_fluxo(ordens, lambda o: o["parceiro"]))

    ws.column_dimensions["A"].width = 48
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = 16


def _media_dias(pares):
    dias = [(fim - ini).total_seconds() / 86400 for ini, fim in pares if ini and fim and fim >= ini]
    return round(sum(dias) / len(dias), 1) if dias else None


def _quebra_conferente(ordens) -> list[list]:
    grupos = defaultdict(lambda: defaultdict(int))
    for o in ordens:
        if o["conferida_no_periodo"]:
            g = grupos[o["conferente"] or _SEM_INFO]
            g["conferidas"] += 1
            if o["como"] == _SEM_CONTAGEM:
                g["sem_contagem"] += 1
            else:
                g["linhas"] += o["linhas_item"]
                g["pecas"] += o["pecas"] or 0
        for t in o["tentativas_periodo"]:
            grupos[t["usuario"] or _SEM_INFO]["tentativas"] += 1
    linhas = [[k, g["conferidas"], g["sem_contagem"], g["linhas"], g["pecas"], g["tentativas"]] for k, g in grupos.items()]
    return sorted(linhas, key=lambda r: (-r[1], -r[5], r[0]))


def _quebra_fluxo(ordens, chave) -> list[list]:
    grupos = defaultdict(lambda: defaultdict(int))
    for o in ordens:
        g = grupos[chave(o) or _SEM_INFO]
        g["conferidas"] += o["conferida_no_periodo"]
        g["faturadas"] += o["faturada_no_periodo"]
        g["expedidas"] += o["expedida_no_periodo"]
        g["divergentes"] += o["divergente"]
    linhas = [[k, g["conferidas"], g["faturadas"], g["expedidas"], g["divergentes"]] for k, g in grupos.items()]
    return sorted(linhas, key=lambda r: (-(r[1] + r[2] + r[3]), r[0]))


def _aba_ordens(ws, ordens):
    colunas = [
        ("Tipo", "tipo", None), ("ID interno", "codigo_interno", None), ("Ordem", "ordem", None),
        ("Cliente / fornecedor", "parceiro", None), ("Orçamento", "orcamento", None), ("Pedido", "pedido", None),
        ("OS", "os", None), ("NF", "nf", None), ("Status atual", "status", None), ("Operação", "operacao", None),
        ("Solicitada em", "solicitada_em", _FMT_DATA), ("Previsão de entrega", "previsao", _FMT_DATA),
        ("Conferente", "conferente", None), ("Conferida em", "conferida_em", _FMT_DATA), ("Como foi conferida", "como", None),
        ("Conferência x faturamento", "conf_x_fat", None), ("Edições depois de conferida", "edicoes", None),
        ("Tentativas bloqueadas", "qtd_tentativas", None), ("Conferido diferente da NF", "nf_div", None),
        ("Faturada em", "faturada_em", _FMT_DATA), ("Expedida em", "expedida_em", _FMT_DATA),
        ("Expedida por", "expedida_por", None), ("Expedida sem conferência", "sem_conf", None),
        ("Linhas de item", "linhas_item", None), ("Peças contadas", "pecas", None), ("Volumes", "volumes", None),
        ("Espécie", "especie", None), ("Peso bruto", "peso_bruto", None), ("Peso líquido", "peso_liquido", None),
    ]
    _lista(ws, colunas, [
        {**o, "qtd_tentativas": len(o["tentativas"]), "nf_div": "Sim" if o["nf_divergente"] else "",
         "sem_conf": "Sim" if o["expedida_sem_conf"] else ""}
        for o in ordens
    ])


def _aba_divergencias(ws, linhas):
    colunas = [
        ("Data", "data", _FMT_DATA), ("Tipo", "tipo", None), ("Ordem", "ordem", None),
        ("Cliente / fornecedor", "parceiro", None), ("Divergência", "divergencia", None), ("Quem", "quem", None),
        ("Código", "codigo", None), ("Item", "item", None), ("OS", "os", None),
        ("Qtd informada / conferida", "qtd", None), ("Qtd na NF", "qtd_nf", None),
    ]
    _lista(ws, colunas, linhas)


def _aba_sem_conferencia(ws, linhas):
    colunas = [
        ("Data", "data", _FMT_DATA), ("Situação", "situacao", None), ("Tipo", "tipo", None), ("Ordem", "ordem", None),
        ("Cliente / fornecedor", "parceiro", None), ("NF", "nf", None), ("Quem", "quem", None), ("Motivo", "motivo", None),
    ]
    _lista(ws, colunas, linhas)


def _descrever_filtros(f) -> str:
    partes = []
    if f["tipo"]:
        partes.append(f"só {TIPOS[f['tipo']]}")
    if f["conferente"]:
        partes.append(f"conferente {f['conferente']}")
    if f["parceiro"]:
        partes.append(f"cliente/fornecedor contém \"{f['parceiro']}\"")
    if f["status"]:
        partes.append(f"status atual {STATUS[f['status']]}")
    if f["so_divergentes"]:
        partes.append("só com divergência")
    return ", ".join(partes) or "nenhum (tudo)"
