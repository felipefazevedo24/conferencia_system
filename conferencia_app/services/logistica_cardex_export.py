"""Exportação do Cardex: Excel no layout do Registro de Inventário Modelo 7
(o mesmo da planilha que a Contabilidade monta à mão) e PDF para arquivar
o fechamento. Recebe o que o logistica_cardex_service já calculou."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .logistica_cardex_service import CLASSE_ROTULO

FMT_QTDE = '#,##0.000;[Red]-#,##0.000'
FMT_VALOR = '#,##0.00;[Red]-#,##0.00'
FMT_CUSTO = '#,##0.0000;[Red]-#,##0.0000'
AZUL = "1B3B6F"
AZUL_CLARO = "8DB4E2"
CINZA = "F2F2F2"
_borda = Side(style="thin", color="BFBFBF")
BORDA = Border(left=_borda, right=_borda, top=_borda, bottom=_borda)


def _data_br(iso: str | None) -> str:
    if not iso:
        return ""
    return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}"


def rotulo_fonte(cab: dict) -> str:
    if cab.get("fonte") == "congelado":
        f = cab.get("fechamento") or {}
        return f"Mês fechado (congelado em {_data_br((f.get('fechado_em') or '')[:10])} por {f.get('fechado_por') or '-'})"
    if cab.get("fonte") == "ancora":
        a = cab.get("ancora") or {}
        return f"Prévia calculada do GRV, abertura pelo fechamento de {a.get('rotulo') or '-'}"
    return "Prévia calculada do GRV (sem mês fechado anterior)"


def _rotulo_depositos(cab: dict) -> str:
    selecionados = cab.get("depositos_selecionados") or []
    if not selecionados:
        return "Todos"
    nomes = {d["codigo"]: d["nome"] for d in cab.get("depositos") or []}
    return ", ".join(f"{c} - {nomes.get(c, '')}".strip(" -") for c in selecionados)


def _cabecalho_planilha(ws, titulo: str, cab: dict, colunas: int) -> int:
    ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=13, color=AZUL)
    periodo = cab["periodo"]
    ws.cell(row=2, column=1, value=f"Período: {_data_br(periodo['inicio'])} a {_data_br(periodo['fim'])}")
    ws.cell(row=3, column=1, value=f"Depósitos: {_rotulo_depositos(cab)}")
    ws.cell(row=4, column=1, value=f"Origem: {rotulo_fonte(cab)}")
    for linha in range(2, 5):
        ws.cell(row=linha, column=1).font = Font(size=9, color="555555")
    return 6


def _estilo_titulos(ws, linha: int, titulos: list[str]) -> None:
    for col, titulo in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=col, value=titulo)
        c.font = Font(bold=True, size=9)
        c.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[linha].height = 30


def _larguras(ws, larguras: list[int]) -> None:
    for idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = largura


RESUMO_COLUNAS = [
    ("Família", "familia", None, 26),
    ("Código", "codigo_interno", None, 14),
    ("Descrição", "descricao", None, 40),
    ("Un.", "unidade", None, 6),
    ("Classif. fiscal", "ncm", None, 12),
    ("Saldo inicial qtde", "qtde_inicial", FMT_QTDE, 14),
    ("Saldo inicial R$", "valor_inicial", FMT_VALOR, 15),
    ("Entradas qtde", "qtde_entrada", FMT_QTDE, 13),
    ("Entradas R$", "valor_entrada", FMT_VALOR, 14),
    ("Saídas qtde", "qtde_saida", FMT_QTDE, 13),
    ("Saídas R$", "valor_saida", FMT_VALOR, 14),
    ("Transf. qtde", "qtde_transferencia", FMT_QTDE, 12),
    ("Transf. R$", "valor_transferencia", FMT_VALOR, 13),
    ("Ajustes qtde", "qtde_ajuste", FMT_QTDE, 12),
    ("Ajustes R$", "valor_ajuste", FMT_VALOR, 13),
    ("Reavaliação R$", "valor_reavaliacao", FMT_VALOR, 13),
    ("Saldo final qtde", "qtde_final", FMT_QTDE, 14),
    ("Saldo final R$", "valor_final", FMT_VALOR, 15),
    ("Custo médio final", "custo_medio_final", FMT_CUSTO, 13),
]


def _aba_resumo(ws, resumo: dict) -> None:
    linha = _cabecalho_planilha(ws, "CARDEX — RESUMO DO PERÍODO", resumo, len(RESUMO_COLUNAS))
    _estilo_titulos(ws, linha, [c[0] for c in RESUMO_COLUNAS])
    _larguras(ws, [c[3] for c in RESUMO_COLUNAS])
    ws.freeze_panes = ws.cell(row=linha + 1, column=4)
    linha += 1

    def escrever(valores: dict, negrito=False, fundo=None):
        nonlocal linha
        for col, (_, chave, fmt, _) in enumerate(RESUMO_COLUNAS, start=1):
            c = ws.cell(row=linha, column=col, value=valores.get(chave))
            c.border = BORDA
            c.font = Font(size=9, bold=negrito)
            if fmt:
                c.number_format = fmt
            if fundo:
                c.fill = PatternFill("solid", fgColor=fundo)
        linha += 1

    familia_atual = None
    for item in resumo["linhas"]:
        if familia_atual is not None and item["familia"] != familia_atual:
            escrever({**resumo["subtotais_familia"][familia_atual], "familia": f"Subtotal {familia_atual}", "custo_medio_final": None}, True, CINZA)
        familia_atual = item["familia"]
        escrever(item)
    if familia_atual is not None:
        escrever({**resumo["subtotais_familia"][familia_atual], "familia": f"Subtotal {familia_atual}", "custo_medio_final": None}, True, CINZA)
    escrever({**resumo["totais"], "familia": "TOTAL GERAL", "custo_medio_final": None}, True, AZUL_CLARO)


CARDEX_COLUNAS = [
    ("Data movimentação", 12), ("Tipo movimento", 14), ("Obs.", 48), ("Documento", 26), ("Cod interno", 13),
    ("Item", 34), ("Un.", 5), ("Classificação fiscal", 12), ("Família", 22), ("Grupo", 7), ("Cod dep.", 7),
    ("Quantidade", 13), ("Custo unitário", 12), ("Total", 13), ("ICMS", 8), ("PIS", 8), ("COFINS", 8),
    ("IBS", 8), ("CBS", 8), ("Saldo quantidade", 14), ("Saldo total", 14), ("Custo médio", 11),
]
_FMT_CARDEX = {11: FMT_QTDE, 12: FMT_CUSTO, 13: FMT_VALOR, 14: FMT_CUSTO, 15: FMT_CUSTO, 16: FMT_CUSTO,
               17: FMT_CUSTO, 18: FMT_CUSTO, 19: FMT_QTDE, 20: FMT_VALOR, 21: FMT_CUSTO}


def _aba_cardex(ws, cab: dict, detalhes: list[dict]) -> None:
    linha = _cabecalho_planilha(ws, "REGISTRO DE INVENTÁRIO — MODELO 7 / CARDEX", cab, len(CARDEX_COLUNAS))
    ws.cell(row=5, column=1, value="** operações de saída utilizam o custo médio").font = Font(size=8, italic=True, color="777777")
    _estilo_titulos(ws, linha, [c[0] for c in CARDEX_COLUNAS])
    _larguras(ws, [c[1] for c in CARDEX_COLUNAS])
    ws.freeze_panes = ws.cell(row=linha + 1, column=4)
    linha += 1

    def escrever(valores: list, negrito=False, fundo=None, vermelho: set | None = None):
        nonlocal linha
        for col, valor in enumerate(valores):
            c = ws.cell(row=linha, column=col + 1, value=valor)
            c.border = BORDA
            c.font = Font(size=9, bold=negrito, color="C00000" if vermelho and col in vermelho else None)
            if col in _FMT_CARDEX:
                c.number_format = _FMT_CARDEX[col]
            if fundo:
                c.fill = PatternFill("solid", fgColor=fundo)
        linha += 1

    for det in detalhes:
        base = [det["codigo_interno"], det["descricao"], det["unidade"], det["ncm"], det["familia"], det["grupo"]]
        ini = det["inicial"]
        escrever(
            ["", "Saldo inicial", "", ""] + base + ["", None, None, None, None, None, None, None, None,
                                                  ini["qtde"], ini["valor"], ini["custo_medio"]],
            True, CINZA,
        )
        for l in det["linhas"]:
            imp = {k: v for k, v in (l.get("impostos_unit") or {}).items() if v}
            informativo = l["classe"] == "informativo"
            escrever(
                [_data_br(l["data"][:10]), f"{l['tipo']}-{CLASSE_ROTULO.get(l['classe'], l['tipo_rotulo'])}", l["obs"], l["documento"]]
                + base
                + [l["deposito"], l["qtde"], l["custo_unit"], l["valor"],
                   imp.get("icms"), imp.get("pis"), imp.get("cofins"), imp.get("ibs"), imp.get("cbs"),
                   None if informativo else l["saldo_qtde"], None if informativo else l["saldo_valor"], l["custo_medio"]],
                # Custo pelo médio em vermelho, como a Contabilidade marca.
                vermelho={12} if l.get("origem_custo") != "nf" and not informativo else None,
            )
        fim = det["final"]
        escrever(
            ["", "Saldo final", "", ""] + base + ["", None, None, None, None, None, None, None, None,
                                                fim["qtde"], fim["valor"], fim["custo_medio"]],
            True, CINZA,
        )
        linha += 1


def gerar_xlsx(resumo: dict | None, cab: dict, detalhes: list[dict] | None) -> bytes:
    wb = Workbook()
    ws = wb.active
    if resumo is not None:
        ws.title = "Resumo"
        _aba_resumo(ws, resumo)
        if detalhes:
            _aba_cardex(wb.create_sheet("Cardex"), cab, detalhes)
    else:
        ws.title = "Cardex"
        _aba_cardex(ws, cab, detalhes or [])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

_styles = getSampleStyleSheet()
_P = ParagraphStyle("cardex", parent=_styles["Normal"], fontName="Helvetica", fontSize=6.5, leading=8)
_P_TITULO = ParagraphStyle("cardex_t", parent=_styles["Normal"], fontName="Helvetica-Bold", fontSize=12, textColor=colors.HexColor("#" + AZUL))
_P_SUB = ParagraphStyle("cardex_s", parent=_styles["Normal"], fontName="Helvetica", fontSize=8, textColor=colors.HexColor("#555555"))


def _br(valor: Any, casas: int = 2) -> str:
    if valor is None:
        return ""
    texto = f"{float(valor):,.{casas}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def _doc_pdf(titulo: str, cab: dict, tabela: Table) -> bytes:
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm, title=titulo)
    periodo = cab["periodo"]
    elementos = [
        Paragraph(titulo, _P_TITULO),
        Paragraph(
            f"Período {_data_br(periodo['inicio'])} a {_data_br(periodo['fim'])} · Depósitos: {_rotulo_depositos(cab)} · {rotulo_fonte(cab)}",
            _P_SUB,
        ),
        Spacer(1, 4 * mm),
        tabela,
    ]
    doc.build(elementos)
    return out.getvalue()


def _estilo_tabela(linhas_destaque: list[int]) -> TableStyle:
    estilo = [
        ("FONT", (0, 0), (-1, -1), "Helvetica", 6.5),
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 6.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + AZUL_CLARO)),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFBFBF")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("REPEATROWS", (0, 0), (-1, 0)),
    ]
    for idx in linhas_destaque:
        estilo.append(("BACKGROUND", (0, idx), (-1, idx), colors.HexColor("#" + CINZA)))
        estilo.append(("FONT", (0, idx), (-1, idx), "Helvetica-Bold", 6.5))
    return TableStyle(estilo)


def gerar_pdf_resumo(resumo: dict) -> bytes:
    cabecalho = ["Código", "Descrição", "Un.", "Saldo inicial", "R$ inicial", "Entradas R$", "Saídas R$",
                 "Ajustes R$", "Transf./Reav. R$", "Saldo final", "R$ final", "Médio final"]
    dados = [cabecalho]
    destaque = []

    def linha(v: dict, codigo: str, descricao: str) -> list:
        return [
            codigo, Paragraph(descricao or "", _P), v.get("unidade", ""),
            _br(v.get("qtde_inicial"), 3), _br(v.get("valor_inicial")), _br(v.get("valor_entrada")),
            _br(v.get("valor_saida")), _br(v.get("valor_ajuste")),
            _br((v.get("valor_transferencia") or 0) + (v.get("valor_reavaliacao") or 0)),
            _br(v.get("qtde_final"), 3), _br(v.get("valor_final")),
            _br(v.get("custo_medio_final"), 4) if "custo_medio_final" in v else "",
        ]

    familia_atual = None
    for item in resumo["linhas"]:
        if item["familia"] != familia_atual:
            if familia_atual is not None:
                destaque.append(len(dados))
                dados.append(linha({**resumo["subtotais_familia"][familia_atual], "unidade": ""}, "", f"Subtotal {familia_atual}"))
            familia_atual = item["familia"]
        dados.append(linha(item, item["codigo_interno"], item["descricao"]))
    if familia_atual is not None:
        destaque.append(len(dados))
        dados.append(linha({**resumo["subtotais_familia"][familia_atual], "unidade": ""}, "", f"Subtotal {familia_atual}"))
    destaque.append(len(dados))
    dados.append(linha({**resumo["totais"], "unidade": ""}, "", "TOTAL GERAL"))

    tabela = Table(dados, colWidths=[22 * mm, 70 * mm, 10 * mm] + [19 * mm] * 8 + [16 * mm], repeatRows=1)
    tabela.setStyle(_estilo_tabela(destaque))
    for col in range(3, len(cabecalho)):
        tabela.setStyle(TableStyle([("ALIGN", (col, 0), (col, -1), "RIGHT")]))
    return _doc_pdf("CARDEX — RESUMO DO PERÍODO", resumo, tabela)


def gerar_pdf_item(cab: dict, det: dict) -> bytes:
    cabecalho = ["Data", "Tipo", "Documento / Obs.", "Dep.", "Quantidade", "Custo unit.", "Total",
                 "ICMS", "PIS", "COFINS", "Saldo qtde", "Saldo total", "Custo médio"]
    dados = [cabecalho]
    ini, fim = det["inicial"], det["final"]
    dados.append(["", "Saldo inicial", "", "", "", "", "", "", "", "", _br(ini["qtde"], 3), _br(ini["valor"]), _br(ini["custo_medio"], 4)])
    for l in det["linhas"]:
        imp = {k: v for k, v in (l.get("impostos_unit") or {}).items() if v}
        informativo = l["classe"] == "informativo"
        dados.append([
            _data_br(l["data"][:10]), CLASSE_ROTULO.get(l["classe"], l["tipo_rotulo"]),
            Paragraph(f"<b>{l['documento']}</b><br/>{l['obs']}", _P), l["deposito"],
            _br(l["qtde"], 3), _br(l["custo_unit"], 4), _br(l["valor"]),
            _br(imp.get("icms"), 4), _br(imp.get("pis"), 4), _br(imp.get("cofins"), 4),
            "" if informativo else _br(l["saldo_qtde"], 3), "" if informativo else _br(l["saldo_valor"]), _br(l["custo_medio"], 4),
        ])
    dados.append(["", "Saldo final", "", "", "", "", "", "", "", "", _br(fim["qtde"], 3), _br(fim["valor"]), _br(fim["custo_medio"], 4)])
    tabela = Table(dados, colWidths=[16 * mm, 22 * mm, 82 * mm, 10 * mm] + [16 * mm] * 9, repeatRows=1)
    tabela.setStyle(_estilo_tabela([1, len(dados) - 1]))
    for col in range(4, len(cabecalho)):
        tabela.setStyle(TableStyle([("ALIGN", (col, 0), (col, -1), "RIGHT")]))
    titulo = f"CARDEX — {det['codigo_interno']} {det['descricao']} ({det['unidade']})"
    return _doc_pdf(titulo, cab, tabela)
