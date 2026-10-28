"""Relatório da Central de Viagens em Excel: o que foi feito num período.

A data de "feito" vem de onde o sistema grava a execução de verdade:
- coleta/entrega dentro de viagem: a parada concluída (saida_real). A
  solicitação concluída pela viagem NÃO ganha concluido_em (ver
  _sync_solicitacao_viagem), então a data dela não serve;
- fluxo antigo, sem viagem: a solicitação concluída (concluido_em);
- viagem: concluída, pela data de retorno.

Centro de custo = Departamento escolhido na alocação; não há campo próprio.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, time, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, exists, func

from ..extensions import db
from ..models import AgendamentoMotorista, AgendamentoSolicitacao, AgendamentoVeiculo, Viagem, ViagemParada
from ..tempo import agora_br

# Mesma lista da alocação (agendamento_routes, departamentos_validos).
DEPARTAMENTOS = ("COMPRAS", "ASSISTÊNCIA TÉCNICA", "ENGENHARIA/PCP", "LOGÍSTICA", "FACILITIES")
TIPOS = {"COLETA": "Coleta", "ENTREGA": "Entrega", "AVULSA": "Uso avulso"}
_ORIGENS = {
    "ORDEM_DE_COMPRA": "Ordem de Compra",
    "SOLICITACAO_USUARIO": "Solicitação pela tela",
    "AutoCIF": "Automática (CIF)",
    "AutoDAP": "Automática (DAP)",
    "ROMANEIO": "Romaneio",
    "Manual": "Manual",
    "GoogleSheets": "Google Sheets",
    "Consyste": "Consyste",
    "ERPPostgres": "ERP",
}
_SEM_INFO = "Não informado"

_FMT_DATA = "dd/mm/yyyy hh:mm"
_FMT_KM = "#,##0.0"
_FMT_MOEDA = '"R$" #,##0.00'
_FMT_HORAS = "#,##0.0"
_CABECALHO_FONTE = Font(bold=True, color="FFFFFF")
_CABECALHO_FUNDO = PatternFill("solid", fgColor="1F3864")
_SECAO_FONTE = Font(bold=True, size=12)


def opcoes() -> dict:
    """Listas do painel. Inclui inativos: viagem antiga pode ter usado um."""
    veiculos = AgendamentoVeiculo.query.order_by(
        AgendamentoVeiculo.ativo.desc(), AgendamentoVeiculo.ordem_exibicao, AgendamentoVeiculo.nome_exibicao
    ).all()
    motoristas = AgendamentoMotorista.query.order_by(AgendamentoMotorista.ativo.desc(), AgendamentoMotorista.nome).all()
    return {
        "veiculos": [{"id": v.id, "nome": _rotulo_veiculo(v), "ativo": bool(v.ativo)} for v in veiculos],
        "motoristas": [{"id": m.id, "nome": m.nome, "ativo": bool(m.ativo)} for m in motoristas],
        "departamentos": list(DEPARTAMENTOS),
        "tipos": [{"valor": k, "nome": v} for k, v in TIPOS.items()],
    }


def montar_filtros(args) -> dict:
    inicio = _data(args.get("inicio"), "Informe a data inicial.")
    fim = _data(args.get("fim"), "Informe a data final.")
    if fim < inicio:
        raise ValueError("A data final é anterior à inicial.")
    departamento = str(args.get("departamento") or "").strip()
    if departamento and departamento not in DEPARTAMENTOS:
        raise ValueError("Departamento inválido.")
    tipo = str(args.get("tipo") or "").strip().upper()
    if tipo and tipo not in TIPOS:
        raise ValueError("Tipo inválido.")
    return {
        "inicio": inicio,
        "fim": fim,
        "veiculo_id": _id(args.get("veiculo_id"), "Veículo inválido."),
        "motorista_id": _id(args.get("motorista_id"), "Motorista inválido."),
        "departamento": departamento,
        "tipo": tipo,
    }


def gerar_excel(args, usuario: str) -> tuple[BytesIO, str]:
    filtros = montar_filtros(args)
    dados = coletar(filtros)
    wb = Workbook()
    _aba_resumo(wb.active, filtros, dados, usuario)
    _aba_viagens(wb.create_sheet("Viagens"), dados["viagens"])
    _aba_atendimentos(wb.create_sheet("Coletas e entregas"), dados["realizados"])
    _aba_nao_feitas(wb.create_sheet("Não feitas"), dados["nao_realizados"] + dados["canceladas"])
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    nome = f"relatorio_viagens_{filtros['inicio']:%Y%m%d}_{filtros['fim']:%Y%m%d}.xlsx"
    return stream, nome


# --------------------------------------------------------------------------- dados

def coletar(f: dict) -> dict:
    ini = datetime.combine(f["inicio"], time.min)
    fim = datetime.combine(f["fim"] + timedelta(days=1), time.min)
    veiculos = {v.id: v for v in AgendamentoVeiculo.query.all()}
    motoristas = {m.id: m.nome for m in AgendamentoMotorista.query.all()}
    ctx = {"veiculos": veiculos, "motoristas": motoristas}

    realizados = _por_parada(f, ini, fim, "Concluida", ctx) + _sem_viagem(f, ini, fim, ctx)
    realizados.sort(key=lambda a: a["data"] or datetime.min)
    nao_realizados = _por_parada(f, ini, fim, "Nao_realizada", ctx)
    canceladas = _canceladas(f, ini, fim, ctx)
    viagens = _viagens(f, ini, fim, ctx)
    return {"viagens": viagens, "realizados": realizados, "nao_realizados": nao_realizados, "canceladas": canceladas}


def _por_parada(f, ini, fim, status, ctx) -> list[dict]:
    data_parada = func.coalesce(ViagemParada.saida_real, ViagemParada.chegada_real, Viagem.retorno_real)
    q = (
        db.session.query(ViagemParada, Viagem, AgendamentoSolicitacao)
        .join(Viagem, Viagem.id == ViagemParada.viagem_id)
        .outerjoin(AgendamentoSolicitacao, AgendamentoSolicitacao.id == ViagemParada.solicitacao_id)
        .filter(
            ViagemParada.status == status,
            ViagemParada.tipo.in_(list(TIPOS)),
            data_parada >= ini,
            data_parada < fim,
        )
    )
    if f["veiculo_id"]:
        q = q.filter(Viagem.veiculo_id == f["veiculo_id"])
    if f["motorista_id"]:
        q = q.filter(Viagem.motorista_id == f["motorista_id"])
    if f["departamento"]:
        q = q.filter(AgendamentoSolicitacao.departamento_solicitante == f["departamento"])
    if f["tipo"]:
        q = q.filter(ViagemParada.tipo == f["tipo"])
    linhas = []
    for parada, viagem, sol in q.all():
        linha = _linha_solicitacao(sol, ctx, viagem.veiculo_id, viagem.motorista_id, viagem.motorista_nome)
        linha.update({
            "data": parada.saida_real or parada.chegada_real or viagem.retorno_real,
            "tipo": str(parada.tipo or "").upper(),
            "viagem": viagem.codigo or f"#{viagem.id}",
            "resultado": parada.resultado or "",
            "observacao": parada.observacao or (sol.observacoes_logistica if sol else "") or "",
            "situacao": "Não realizada" if status == "Nao_realizada" else "Feita",
            "registrado_por": "",
        })
        if not sol:
            # Parada colocada direto na viagem, sem solicitação por trás.
            linha.update({
                "parceiro": parada.parceiro_nome or "",
                "endereco": parada.endereco or "",
                "cidade": parada.cidade or "",
                "uf": parada.uf or "",
                "origem": "Parada incluída na viagem",
            })
        linhas.append(linha)
    return linhas


def _sem_viagem(f, ini, fim, ctx) -> list[dict]:
    """Fluxo antigo: concluída direto na solicitação. Fica de fora a que tem
    parada concluída, que já entrou pela viagem."""
    tem_parada = exists().where(and_(
        ViagemParada.solicitacao_id == AgendamentoSolicitacao.id,
        ViagemParada.status == "Concluida",
    ))
    q = AgendamentoSolicitacao.query.filter(
        AgendamentoSolicitacao.status == "Concluida",
        AgendamentoSolicitacao.concluido_em >= ini,
        AgendamentoSolicitacao.concluido_em < fim,
        ~tem_parada,
    )
    linhas = []
    for sol in _filtrar_solicitacoes(q, f).all():
        linha = _linha_solicitacao(sol, ctx, sol.veiculo_id, sol.motorista_id, sol.motorista_nome)
        linha.update({
            "data": sol.concluido_em,
            "tipo": str(sol.tipo or "").upper(),
            "viagem": "Sem viagem",
            "resultado": "",
            "observacao": sol.observacoes_logistica or "",
        })
        linhas.append(linha)
    return linhas


def _canceladas(f, ini, fim, ctx) -> list[dict]:
    q = AgendamentoSolicitacao.query.filter(
        AgendamentoSolicitacao.status == "Cancelada",
        AgendamentoSolicitacao.cancelado_em >= ini,
        AgendamentoSolicitacao.cancelado_em < fim,
    )
    linhas = []
    for sol in _filtrar_solicitacoes(q, f).order_by(AgendamentoSolicitacao.cancelado_em).all():
        linha = _linha_solicitacao(sol, ctx, sol.veiculo_id, sol.motorista_id, sol.motorista_nome)
        linha.update({
            "data": sol.cancelado_em,
            "tipo": str(sol.tipo or "").upper(),
            "viagem": "",
            "resultado": sol.motivo_cancelamento or "",
            "observacao": "",
            "situacao": "Cancelada",
            "registrado_por": sol.cancelado_por or "",
        })
        linhas.append(linha)
    return linhas


def _filtrar_solicitacoes(q, f):
    if f["veiculo_id"]:
        q = q.filter(AgendamentoSolicitacao.veiculo_id == f["veiculo_id"])
    if f["motorista_id"]:
        q = q.filter(AgendamentoSolicitacao.motorista_id == f["motorista_id"])
    if f["departamento"]:
        q = q.filter(AgendamentoSolicitacao.departamento_solicitante == f["departamento"])
    if f["tipo"]:
        q = q.filter(AgendamentoSolicitacao.tipo == f["tipo"])
    return q


def _linha_solicitacao(sol, ctx, veiculo_id, motorista_id, motorista_nome) -> dict:
    veiculo = ctx["veiculos"].get(veiculo_id)
    linha = {
        "veiculo": _rotulo_veiculo(veiculo) if veiculo else "",
        "motorista": motorista_nome or ctx["motoristas"].get(motorista_id) or "",
        "solicitacao": "", "documento": "", "orcamento": "", "parceiro": "", "endereco": "",
        "cidade": "", "uf": "", "origem": "", "solicitante": "", "departamento": "",
        "pedida_em": None, "prioridade": "", "alocado_por": "", "itens": "",
    }
    if sol:
        documento = " ".join(filter(None, [sol.documento_tipo, sol.documento_numero]))
        linha.update({
            "solicitacao": sol.codigo or f"#{sol.id}",
            "documento": documento,
            "orcamento": sol.orcamento or "",
            "parceiro": sol.parceiro_nome or "",
            "endereco": ", ".join(filter(None, [sol.logradouro, sol.numero, sol.bairro])),
            "cidade": sol.cidade or "",
            "uf": sol.uf or "",
            "origem": _ORIGENS.get(str(sol.origem_documento or "").strip(), sol.origem_documento or ""),
            "solicitante": sol.solicitante or "",
            "departamento": sol.departamento_solicitante or "",
            "pedida_em": sol.criado_em,
            "prioridade": sol.prioridade or "",
            "alocado_por": sol.alocado_por or "",
            "itens": sol.resumo_itens or "",
        })
    return linha


def _viagens(f, ini, fim, ctx) -> list[dict]:
    data_viagem = func.coalesce(Viagem.retorno_real, Viagem.saida_real, Viagem.atualizado_em)
    q = Viagem.query.filter(Viagem.status == "Concluida", data_viagem >= ini, data_viagem < fim)
    if f["veiculo_id"]:
        q = q.filter(Viagem.veiculo_id == f["veiculo_id"])
    if f["motorista_id"]:
        q = q.filter(Viagem.motorista_id == f["motorista_id"])
    viagens = q.order_by(data_viagem).all()

    contagem = defaultdict(lambda: defaultdict(int))
    casa_filtro = set()
    if viagens:
        paradas = (
            db.session.query(ViagemParada.viagem_id, ViagemParada.tipo, ViagemParada.status,
                             AgendamentoSolicitacao.departamento_solicitante)
            .outerjoin(AgendamentoSolicitacao, AgendamentoSolicitacao.id == ViagemParada.solicitacao_id)
            .filter(ViagemParada.viagem_id.in_([v.id for v in viagens]))
            .all()
        )
        for viagem_id, tipo, status, departamento in paradas:
            tipo = str(tipo or "").upper()
            if status == "Concluida" and tipo in TIPOS:
                contagem[viagem_id][tipo] += 1
                if (not f["tipo"] or tipo == f["tipo"]) and (not f["departamento"] or departamento == f["departamento"]):
                    casa_filtro.add(viagem_id)
            elif status == "Nao_realizada":
                contagem[viagem_id]["NAO"] += 1

    linhas = []
    for v in viagens:
        # Com filtro de tipo/departamento, entra a viagem que fez pelo menos
        # uma coleta/entrega desse filtro. Viagem avulsa conta como uso avulso.
        if f["tipo"] or f["departamento"]:
            avulsa_ok = v.avulsa and f["tipo"] == "AVULSA" and not f["departamento"]
            if v.id not in casa_filtro and not avulsa_ok:
                continue
        veiculo = ctx["veiculos"].get(v.veiculo_id)
        minutos = v.tempo_total_min or (
            int((v.retorno_real - v.saida_real).total_seconds() // 60) if v.retorno_real and v.saida_real else 0
        )
        linhas.append({
            "codigo": v.codigo or f"#{v.id}",
            "status": v.status,
            "tipo": v.tipo or "",
            "avulsa": bool(v.avulsa),
            "veiculo": _rotulo_veiculo(veiculo) if veiculo else "",
            "placa": (veiculo.placa or "") if veiculo else "",
            "motorista": v.motorista_nome or ctx["motoristas"].get(v.motorista_id) or "",
            "saida": v.saida_real,
            "retorno": v.retorno_real,
            "horas": round(max(minutos, 0) / 60, 1),
            "km": float(v.km_percorrido or 0),
            "litros": float(v.total_litros or 0),
            "gasto": float(v.total_gasto or 0),
            "coletas": contagem[v.id]["COLETA"],
            "entregas": contagem[v.id]["ENTREGA"],
            "avulsos": contagem[v.id]["AVULSA"],
            "nao_realizadas": contagem[v.id]["NAO"],
            "origem": v.origem_label or "",
            "destino": v.destino_label or "",
            "criado_por": v.criado_por or "",
            "liberada_por": v.liberada_por or "",
            "concluido_por": v.concluido_por or "",
            "responsavel": v.funcionario_responsavel or "",
            "observacao": v.observacao or "",
        })
    return linhas


# --------------------------------------------------------------------------- planilha

def _aba_resumo(ws, f, dados, usuario):
    ws.title = "Resumo"
    viagens, realizados = dados["viagens"], dados["realizados"]
    por_tipo = defaultdict(int)
    for a in realizados:
        por_tipo[a["tipo"]] += 1

    ws.append(["Relatório da Central de Viagens"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Período: {f['inicio']:%d/%m/%Y} a {f['fim']:%d/%m/%Y}"])
    ws.append([f"Filtros: {_descrever_filtros(f)}"])
    ws.append([f"Gerado em {agora_br():%d/%m/%Y %H:%M} por {usuario}"])
    ws.append([])

    _secao(ws, "Números do período", ["", "Total"])
    numeros = [
        ("Viagens realizadas", len(viagens), None),
        ("   das quais avulsas (sem paradas)", sum(1 for v in viagens if v["avulsa"]), None),
        ("Coletas realizadas", por_tipo["COLETA"], None),
        ("Entregas realizadas", por_tipo["ENTREGA"], None),
        ("Usos avulsos realizados", por_tipo["AVULSA"], None),
        ("Total de coletas + entregas + usos avulsos", len(realizados), None),
        ("Tentativas não realizadas", len(dados["nao_realizados"]), None),
        ("Solicitações canceladas", len(dados["canceladas"]), None),
        ("KM rodados", sum(v["km"] for v in viagens), _FMT_KM),
        ("Horas em viagem", sum(v["horas"] for v in viagens), _FMT_HORAS),
        ("Litros abastecidos", sum(v["litros"] for v in viagens), _FMT_KM),
        ("Gasto com abastecimento", sum(v["gasto"] for v in viagens), _FMT_MOEDA),
    ]
    for rotulo, valor, formato in numeros:
        _append(ws, [rotulo, valor])
        if formato:
            ws.cell(row=ws.max_row, column=2).number_format = formato
    ws.append([])

    frota_cab = ["", "Viagens", "KM", "Coletas", "Entregas", "Usos avulsos", "Não realizadas"]
    _tabela(ws, "Por veículo", frota_cab, _quebra_frota(dados, "veiculo"), {3: _FMT_KM})
    _tabela(ws, "Por motorista", frota_cab, _quebra_frota(dados, "motorista"), {3: _FMT_KM})
    tipo_cab = ["", "Coletas", "Entregas", "Usos avulsos", "Total"]
    _tabela(ws, "Por departamento (centro de custo)", tipo_cab, _quebra(realizados, lambda a: a["departamento"]))
    _tabela(ws, "Por origem (como entrou no sistema)", tipo_cab, _quebra(realizados, lambda a: a["origem"]))
    _tabela(ws, "Por solicitante (quem colocou)", tipo_cab, _quebra(realizados, lambda a: a["solicitante"]))
    _tabela(ws, "Por cidade", tipo_cab, _quebra(realizados, lambda a: "/".join(filter(None, [a["cidade"], a["uf"]]))))

    ws.column_dimensions["A"].width = 46
    for col in range(2, 8):
        ws.column_dimensions[get_column_letter(col)].width = 15


def _quebra(atendimentos, chave) -> list[list]:
    grupos = defaultdict(lambda: defaultdict(int))
    for a in atendimentos:
        grupos[chave(a) or _SEM_INFO][a["tipo"]] += 1
    linhas = [[k, c["COLETA"], c["ENTREGA"], c["AVULSA"], sum(c.values())] for k, c in grupos.items()]
    return sorted(linhas, key=lambda r: (-r[-1], r[0]))


def _quebra_frota(dados, campo) -> list[list]:
    grupos = defaultdict(lambda: defaultdict(float))
    for v in dados["viagens"]:
        g = grupos[v[campo] or _SEM_INFO]
        g["viagens"] += 1
        g["km"] += v["km"]
    for a in dados["realizados"]:
        grupos[a[campo] or _SEM_INFO][a["tipo"]] += 1
    for a in dados["nao_realizados"]:
        grupos[a[campo] or _SEM_INFO]["NAO"] += 1
    linhas = [
        [k, int(g["viagens"]), g["km"], int(g["COLETA"]), int(g["ENTREGA"]), int(g["AVULSA"]), int(g["NAO"])]
        for k, g in grupos.items()
    ]
    return sorted(linhas, key=lambda r: (-r[1], -(r[3] + r[4] + r[5]), r[0]))


def _aba_viagens(ws, viagens):
    colunas = [
        ("Viagem", "codigo", None), ("Tipo", "tipo", None), ("Avulsa", "avulsa", None),
        ("Veículo", "veiculo", None), ("Placa", "placa", None), ("Motorista", "motorista", None),
        ("Saída", "saida", _FMT_DATA), ("Retorno", "retorno", _FMT_DATA), ("Horas", "horas", _FMT_HORAS),
        ("KM", "km", _FMT_KM), ("Litros", "litros", _FMT_KM), ("Gasto (R$)", "gasto", _FMT_MOEDA),
        ("Coletas", "coletas", None), ("Entregas", "entregas", None), ("Usos avulsos", "avulsos", None),
        ("Não realizadas", "nao_realizadas", None), ("Origem", "origem", None), ("Destino", "destino", None),
        ("Criada por", "criado_por", None), ("Liberada por", "liberada_por", None),
        ("Concluída por", "concluido_por", None), ("Responsável (avulsa)", "responsavel", None),
        ("Observação", "observacao", None),
    ]
    linhas = [{**v, "avulsa": "Sim" if v["avulsa"] else "Não"} for v in viagens]
    _lista(ws, colunas, linhas)


def _aba_atendimentos(ws, realizados):
    colunas = [
        ("Feita em", "data", _FMT_DATA), ("Tipo", "tipo_nome", None), ("Solicitação", "solicitacao", None),
        ("Documento", "documento", None), ("Orçamento", "orcamento", None), ("Parceiro", "parceiro", None),
        ("Endereço", "endereco", None), ("Cidade", "cidade", None), ("UF", "uf", None),
        ("Origem (como entrou)", "origem", None), ("Solicitante (quem colocou)", "solicitante", None),
        ("Departamento", "departamento", None), ("Pedida em", "pedida_em", _FMT_DATA),
        ("Prioridade", "prioridade", None), ("Viagem", "viagem", None), ("Veículo", "veiculo", None),
        ("Motorista", "motorista", None), ("Resultado", "resultado", None), ("Alocada por", "alocado_por", None),
        ("Itens / volumes", "itens", None), ("Observação", "observacao", None),
    ]
    _lista(ws, colunas, [{**a, "tipo_nome": TIPOS.get(a["tipo"], a["tipo"])} for a in realizados])


def _aba_nao_feitas(ws, linhas):
    colunas = [
        ("Data", "data", _FMT_DATA), ("Situação", "situacao", None), ("Tipo", "tipo_nome", None),
        ("Solicitação", "solicitacao", None), ("Documento", "documento", None), ("Parceiro", "parceiro", None),
        ("Cidade", "cidade", None), ("UF", "uf", None), ("Solicitante (quem colocou)", "solicitante", None),
        ("Departamento", "departamento", None), ("Viagem", "viagem", None), ("Veículo", "veiculo", None),
        ("Motorista", "motorista", None), ("Motivo / resultado", "motivo", None),
        ("Cancelada por", "registrado_por", None),
    ]
    linhas = sorted(linhas, key=lambda a: a["data"] or datetime.min)
    _lista(ws, colunas, [
        {**a, "tipo_nome": TIPOS.get(a["tipo"], a["tipo"]),
         "motivo": " - ".join(filter(None, [a["resultado"], a["observacao"]]))}
        for a in linhas
    ])


def _secao(ws, titulo, cabecalho):
    ws.append([titulo])
    ws.cell(row=ws.max_row, column=1).font = _SECAO_FONTE
    ws.append(cabecalho)
    for col in range(1, len(cabecalho) + 1):
        cel = ws.cell(row=ws.max_row, column=col)
        cel.font = _CABECALHO_FONTE
        cel.fill = _CABECALHO_FUNDO


def _tabela(ws, titulo, cabecalho, linhas, formatos=None):
    _secao(ws, titulo, cabecalho)
    if not linhas:
        ws.append(["Nada no período"])
    for linha in linhas:
        _append(ws, linha)
        for col, formato in (formatos or {}).items():
            ws.cell(row=ws.max_row, column=col).number_format = formato
    ws.append([])


def _lista(ws, colunas, linhas):
    ws.append([c[0] for c in colunas])
    for col in range(1, len(colunas) + 1):
        cel = ws.cell(row=1, column=col)
        cel.font = _CABECALHO_FONTE
        cel.fill = _CABECALHO_FUNDO
    larguras = [len(c[0]) for c in colunas]
    for linha in linhas:
        _append(ws, [linha.get(chave) if linha.get(chave) is not None else "" for _, chave, _ in colunas])
        for idx, (_, chave, formato) in enumerate(colunas):
            valor = linha.get(chave)
            if formato and valor not in (None, ""):
                ws.cell(row=ws.max_row, column=idx + 1).number_format = formato
            larguras[idx] = max(larguras[idx], 16 if isinstance(valor, datetime) else len(str(valor or "")))
    for idx, largura in enumerate(larguras):
        ws.column_dimensions[get_column_letter(idx + 1)].width = min(largura + 2, 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# --------------------------------------------------------------------------- apoio

def _append(ws, valores):
    """Texto digitado (observação, parceiro...) que começa com "=" o openpyxl
    grava como fórmula. Força texto."""
    ws.append(valores)
    for cel in ws[ws.max_row]:
        if cel.data_type == "f":
            cel.data_type = "s"


def _descrever_filtros(f) -> str:
    partes = []
    if f["veiculo_id"]:
        veiculo = db.session.get(AgendamentoVeiculo, f["veiculo_id"])
        partes.append(f"veículo {_rotulo_veiculo(veiculo) if veiculo else f['veiculo_id']}")
    if f["motorista_id"]:
        motorista = db.session.get(AgendamentoMotorista, f["motorista_id"])
        partes.append(f"motorista {motorista.nome if motorista else f['motorista_id']}")
    if f["departamento"]:
        partes.append(f"departamento {f['departamento']}")
    if f["tipo"]:
        partes.append(f"tipo {TIPOS[f['tipo']]}")
    return ", ".join(partes) or "nenhum (tudo)"


def _rotulo_veiculo(v) -> str:
    nome = v.nome_exibicao or v.codigo
    return f"{nome} ({v.placa})" if v.placa else nome


def _data(valor, msg: str) -> date:
    texto = str(valor or "").strip()
    if not texto:
        raise ValueError(msg)
    try:
        return date.fromisoformat(texto)
    except ValueError:
        raise ValueError("Data inválida.") from None


def _id(valor, msg: str) -> int | None:
    texto = str(valor or "").strip()
    if not texto:
        return None
    if not texto.isdigit():
        raise ValueError(msg)
    return int(texto)
