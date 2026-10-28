"""Relatório do Documento de Entrada em Excel: notas do Sync + compras do GRV.

Notas: cada número usa a data do próprio evento (entrada no Sync, auditoria,
conferência física, lançamento, estorno). Uma nota pode entrar em mais de um
número.

Compras: OC pela data dela no GRV (tord_com.data). Comprador = Solicitante da
OC, o mesmo campo que o Comex usa. GRV fora não derruba o relatório: as notas
saem e a planilha avisa que as compras ficaram de fora - a não ser com filtro
de comprador, que sem o GRV não tem como aplicar.

Itens: os das OCs da aba de OCs (mesmos filtros), com custo, família/grupo e
quem pediu na solicitação de compra. Custo somado sem item cancelado e sem OC
cancelada.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from io import BytesIO

from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_

from ..compras.services import compras_service
from ..extensions import db
from ..models import ItemNota, LogEstornoLancamento, LogManifestacaoDestinatario, LogTentativaConferencia
from ..tempo import agora_br
from .viagem_relatorio_service import _FMT_DATA, _FMT_MOEDA, _SEM_INFO, _append, _data, _lista, _secao, _tabela

TIPOS = {"nfe": "NF-e", "nfse": "NFS-e", "nao_fiscal": "Não fiscal"}
ETAPAS = {"auditoria": "Auditoria", "lancamento": "Lançamento", "finalizado": "Finalizado"}
_LANCADO = {"Lançado", "Lancado", "LanÃ§ado"}
_RECUSADO = {"xml recusado", "recusado"}
_SEM_OC = "Sem OC"
_SEM_SOLICITACAO = "Sem solicitação (direto na OC)"
_FMT_DIAS = "#,##0.0"
# Preço unitário tem item de R$ 0,0855: duas casas arredondariam o custo.
_FMT_UNITARIO = '"R$" #,##0.00##'
_FMT_DIA = "dd/mm/yyyy"


def montar_filtros(args) -> dict:
    inicio = _data(args.get("inicio"), "Informe a data inicial.")
    fim = _data(args.get("fim"), "Informe a data final.")
    if fim < inicio:
        raise ValueError("A data final é anterior à inicial.")
    tipo = str(args.get("tipo") or "").strip().lower()
    if tipo and tipo not in TIPOS:
        raise ValueError("Tipo de documento inválido.")
    etapa = str(args.get("etapa") or "").strip().lower()
    if etapa and etapa not in ETAPAS:
        raise ValueError("Etapa inválida.")
    return {
        "inicio": inicio,
        "fim": fim,
        "tipo": tipo,
        "etapa": etapa,
        "fornecedor": str(args.get("fornecedor") or "").strip()[:100],
        "comprador": str(args.get("comprador") or "").strip()[:100],
    }


def gerar_excel(args, usuario: str) -> tuple[BytesIO, str]:
    filtros = montar_filtros(args)
    dados = coletar(filtros)
    wb = Workbook()
    _aba_resumo(wb.active, filtros, dados, usuario)
    _aba_notas(wb.create_sheet("Notas"), dados["notas"])
    _aba_ocs(wb.create_sheet("Ordens de compra"), dados)
    _aba_itens_oc(wb.create_sheet("Itens das OCs"), dados)
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    nome = f"relatorio_documento_entrada_{filtros['inicio']:%Y%m%d}_{filtros['fim']:%Y%m%d}.xlsx"
    return stream, nome


# --------------------------------------------------------------------------- dados

def coletar(f: dict) -> dict:
    ini = datetime.combine(f["inicio"], time.min)
    fim = datetime.combine(f["fim"] + timedelta(days=1), time.min)

    def no_periodo(valor) -> bool:
        return valor is not None and ini <= valor < fim

    notas = _notas(ini, fim)
    codigos = sorted({int(c) for n in notas for c in n["ocs"] if c.isdigit()})
    ocs, grv_ok = _ocs_grv(f, codigos)
    if not grv_ok and f["comprador"]:
        raise ValueError("O GRV não respondeu agora, e sem ele não dá para filtrar por comprador. Tente de novo ou tire o filtro.")

    resultado = []
    for n in notas:
        n["entrou_no_periodo"] = no_periodo(n["entrou_em"])
        n["auditada_no_periodo"] = n["auditoria"] in ("Aprovada", "Recusada") and no_periodo(n["auditada_em"])
        n["conferida_no_periodo"] = no_periodo(n["conferida_em"])
        n["lancada_no_periodo"] = no_periodo(n["lancada_em"])
        n["estornos_periodo"] = [e for e in n["estornos"] if no_periodo(e)]
        if not (n["entrou_no_periodo"] or n["auditada_no_periodo"] or n["conferida_no_periodo"]
                or n["lancada_no_periodo"] or n["estornos_periodo"]):
            continue
        dados_oc = [ocs[c] for c in n["ocs"] if c in ocs]
        n["compradores"] = sorted({o["comprador"] for o in dados_oc if o["comprador"]}, key=str.casefold)
        n["comprador"] = ", ".join(n["compradores"])
        n["data_oc"] = min((o["data_oc"] for o in dados_oc if o["data_oc"]), default=None)
        if f["tipo"] and n["tipo"] != f["tipo"]:
            continue
        if f["etapa"] and n["etapa"] != f["etapa"]:
            continue
        if f["fornecedor"] and not _contem(n["fornecedor"], f["fornecedor"]):
            continue
        if f["comprador"] and not any(_contem(c, f["comprador"]) for c in n["compradores"]):
            continue
        resultado.append(n)
    resultado.sort(key=lambda n: (n["entrou_em"] or datetime.min, n["numero"]))

    geradas = []
    for o in ocs.values():
        if not no_periodo(o["data_oc"]):
            continue
        if f["fornecedor"] and not _contem(o["fornecedor"], f["fornecedor"]):
            continue
        if f["comprador"] and not _contem(o["comprador"], f["comprador"]):
            continue
        geradas.append(o)
    geradas.sort(key=lambda o: (o["data_oc"], _num(o["codigo"])))
    _notas_das_ocs(geradas)
    itens, itens_ok = _itens_grv(f, geradas) if grv_ok else ([], False)
    return {"notas": resultado, "ocs": geradas, "grv_ok": grv_ok, "itens": itens, "itens_ok": itens_ok}


def _notas(ini: datetime, fim: datetime) -> list[dict]:
    def periodo(col):
        return and_(col >= ini, col < fim)

    numeros = {n for (n,) in db.session.query(ItemNota.numero_nota).filter(or_(
        periodo(ItemNota.data_importacao), periodo(ItemNota.auditor_data),
        periodo(ItemNota.fim_conferencia), periodo(ItemNota.data_lancamento),
    )).distinct()}
    numeros |= {n for (n,) in db.session.query(LogEstornoLancamento.numero_nota).filter(
        periodo(LogEstornoLancamento.data_estorno)).distinct()}
    numeros.discard(None)
    if not numeros:
        return []
    numeros = sorted(numeros)

    grupos = defaultdict(list)
    for item in ItemNota.query.filter(ItemNota.numero_nota.in_(numeros)).order_by(ItemNota.id):
        grupos[(str(item.numero_nota or "").strip(), _ref(item))].append(item)

    # Os logs da nota são gravados só pelo número (é assim no resto do sistema).
    tentativas = defaultdict(lambda: {"max": 0, "divergencia": False})
    for t in LogTentativaConferencia.query.filter(LogTentativaConferencia.numero_nota.in_(numeros)):
        reg = tentativas[str(t.numero_nota).strip()]
        reg["max"] = max(reg["max"], int(t.tentativa_numero or 0))
        reg["divergencia"] |= str(t.status_item or "").upper() == "DIVERGÊNCIA"
    estornos = defaultdict(list)
    for e in LogEstornoLancamento.query.filter(LogEstornoLancamento.numero_nota.in_(numeros)):
        estornos[str(e.numero_nota).strip()].append(e.data_estorno)
    manifestacoes = {}
    for m in LogManifestacaoDestinatario.query.filter(LogManifestacaoDestinatario.numero_nota.in_(numeros)).order_by(
            LogManifestacaoDestinatario.data):
        manifestacoes[str(m.numero_nota).strip()] = f"{m.manifestacao} ({m.status})"

    notas = []
    for (numero, _), itens in grupos.items():
        primeiro = min(itens, key=lambda i: (i.data_importacao or datetime.max, i.id))
        auditado = max((i for i in itens if i.auditor_data), key=lambda i: i.auditor_data, default=None)
        conferido = max((i for i in itens if i.fim_conferencia), key=lambda i: i.fim_conferencia, default=None)
        lancado = max((i for i in itens if i.data_lancamento), key=lambda i: i.data_lancamento, default=None)
        decisoes = {str(i.auditor_decisao or "").strip().lower() for i in itens}
        recusada = bool(decisoes & _RECUSADO)
        tipo_doc = str(primeiro.tipo_documento or "NFE").upper()
        valores_nf = [float(i.valor_nf) for i in itens if i.valor_nf]
        tent = tentativas.get(numero, {"max": 0, "divergencia": False})
        notas.append({
            "tipo": "nfe" if tipo_doc == "NFE" else ("nfse" if tipo_doc == "NFSE" else "nao_fiscal"),
            "tipo_doc": TIPOS["nfe"] if tipo_doc == "NFE" else (TIPOS["nfse"] if tipo_doc == "NFSE" else f"Não fiscal ({tipo_doc.capitalize()})"),
            "numero": numero,
            "fornecedor": next((i.fornecedor.strip() for i in itens if i.fornecedor and i.fornecedor.strip()), ""),
            "cnpj": next((i.cnpj_emitente for i in itens if i.cnpj_emitente), ""),
            "chave": next((i.chave_acesso for i in itens if i.chave_acesso), ""),
            "emissao": min((i.data_emissao for i in itens if i.data_emissao), default=None),
            "entrou_em": primeiro.data_importacao,
            "importada_por": primeiro.usuario_importacao or "",
            "etapa": _etapa(itens, recusada),
            "situacao": _situacao(itens, recusada),
            "auditoria": "Recusada" if recusada else ("Aprovada" if "xml aprovado" in decisoes else "Pendente"),
            "auditada_em": auditado.auditor_data if auditado else None,
            "auditada_por": (auditado.auditor_usuario or "") if auditado else "",
            "conferida_em": conferido.fim_conferencia if conferido else None,
            "conferida_por": (conferido.usuario_conferencia or "") if conferido else "",
            "sem_conferencia": any(i.sem_conferencia_logistica for i in itens),
            "tentativas": tent["max"],
            "divergencia": tent["divergencia"],
            "lancada_em": lancado.data_lancamento if lancado else None,
            "lancada_por": (lancado.usuario_lancamento or "") if lancado else "",
            "numero_lancamento": next((i.numero_lancamento for i in itens if i.numero_lancamento), "") or "",
            "valor": max(valores_nf) if valores_nf else sum(float(i.valor_produto or 0) for i in itens),
            "ocs": _ocs_da_nota(itens),
            "cfops": ", ".join(sorted({str(i.cfop).strip() for i in itens if i.cfop and str(i.cfop).strip()})),
            "material_cliente": any(i.material_cliente for i in itens),
            "remessa": any(i.remessa for i in itens),
            "mercado_livre": any(i.mercado_livre for i in itens),
            "retorno_industrializacao": any(i.retorno_industrializacao for i in itens),
            "manifestacao": manifestacoes.get(numero, ""),
            "estornos": estornos.get(numero, []),
            "itens": itens,
        })
    return notas


def _ocs_grv(f: dict, codigos: list[int]) -> tuple[dict, bool]:
    try:
        linhas = compras_service.ordens_compra_periodo(f["inicio"].isoformat(), f["fim"].isoformat(), codigos)
    except Exception:  # noqa: BLE001 - GRV fora não derruba o relatório
        current_app.logger.warning("Relatorio do Documento de Entrada: GRV indisponivel", exc_info=True)
        return {}, False
    ocs = {}
    for linha in linhas:
        codigo = str(linha.get("cod_ordem_compra") or "").strip()
        if not codigo:
            continue
        qtd_itens = int(linha.get("qtd_itens") or 0)
        com_saldo = int(linha.get("itens_com_saldo") or 0)
        cancelada = linha.get("cancelada") in (True, "true", "t", 1)
        if cancelada:
            situacao = "Cancelada"
        elif not qtd_itens:
            situacao = "Sem itens"
        elif not com_saldo:
            situacao = "Recebida"
        elif com_saldo < qtd_itens:
            situacao = "Recebida em parte"
        else:
            situacao = "Em aberto"
        ocs[codigo] = {
            "codigo": codigo,
            "data_oc": _dt(linha.get("data_oc")),
            "fornecedor": str(linha.get("fornecedor") or "").strip(),
            "cnpj": str(linha.get("cnpj") or "").strip(),
            "comprador": str(linha.get("comprador") or "").strip(),
            "status_oc": str(linha.get("status_oc") or "").strip(),
            "cancelada": cancelada,
            "situacao": situacao,
            "total": _float(linha.get("total")),
            "cond_pagamento": str(linha.get("cond_pagamento") or "").strip(),
            "forma_pgto": str(linha.get("forma_pgto") or "").strip(),
            "qtd_itens": qtd_itens,
            "itens_com_saldo": com_saldo,
        }
    return ocs, True


def _itens_grv(f: dict, ocs: list[dict]) -> tuple[list[dict], bool]:
    """Itens das OCs da aba de OCs: a consulta traz o período inteiro e aqui
    ficam só os das OCs que passaram nos filtros."""
    if not ocs:
        return [], True
    try:
        linhas = compras_service.itens_ordens_compra_periodo(f["inicio"].isoformat(), f["fim"].isoformat())
    except Exception:  # noqa: BLE001 - GRV fora não derruba o relatório
        current_app.logger.warning("Relatorio do Documento de Entrada: itens das OCs indisponiveis", exc_info=True)
        return [], False
    por_codigo = {o["codigo"]: o for o in ocs}
    itens = []
    for linha in linhas:
        oc = por_codigo.get(str(linha.get("cod_ordem_compra") or "").strip())
        if not oc:
            continue
        cancelado = linha.get("cancelado") in (True, "true", "t", 1)
        if cancelado:
            situacao = "Cancelado"
        elif oc["cancelada"]:
            situacao = "OC cancelada"
        elif _float(linha.get("saldo")) <= 0.0001:
            situacao = "Recebido"
        elif _float(linha.get("entregue")) > 0.0001:
            situacao = "Recebido em parte"
        else:
            situacao = "Em aberto"
        itens.append({
            "oc": oc["codigo"],
            "data_oc": oc["data_oc"],
            "fornecedor": oc["fornecedor"],
            "comprador": oc["comprador"],
            "tipo": str(linha.get("tipo") or "").strip(),
            "item": linha.get("item"),
            "codigo": str(linha.get("codigo") or "").strip(),
            "descricao": str(linha.get("descricao") or "").strip(),
            "familia": str(linha.get("familia") or "").strip(),
            "grupo": str(linha.get("grupo") or "").strip(),
            "unidade": str(linha.get("unidade") or "").strip(),
            "qtde": _float(linha.get("qtde")),
            "valor_unitario": _float(linha.get("valor_unitario")),
            "total": _float(linha.get("total")),
            "ipi": _float(linha.get("ipi")),
            "solicitacao": str(linha.get("solicitacao") or "").strip(),
            "solicitante": str(linha.get("solicitante") or "").strip(),
            "solicitada_em": _dt(linha.get("dt_solicitacao")),
            "os": str(linha.get("os") or "").strip(),
            "situacao": situacao,
            "conta": not cancelado and not oc["cancelada"],
        })
    itens.sort(key=lambda i: (i["data_oc"] or datetime.min, _num(i["oc"]), i["tipo"] != "Material", i["item"] or 0))
    return itens, True


def _notas_das_ocs(ocs: list[dict]) -> None:
    """Notas do Sync que citam cada OC, de qualquer data: mostra se a OC
    gerada no período já teve nota."""
    if not ocs:
        return
    por_oc = defaultdict(set)
    codigos = [o["codigo"] for o in ocs]
    for pedido, numero in db.session.query(ItemNota.pedido_compra, ItemNota.numero_nota).filter(
            ItemNota.pedido_compra.in_(codigos)).distinct():
        por_oc[str(pedido).strip()].add(str(numero or "").strip())
    for o in ocs:
        o["notas"] = ", ".join(sorted(por_oc.get(o["codigo"], set()), key=_num))


def _etapa(itens, recusada: bool) -> str:
    """Mesmas regras das filas da tela (api_routes.documento_entrada_lista)."""
    status = {str(i.status or "").strip() for i in itens}
    if recusada or status & _LANCADO:
        return "finalizado"
    if status and status <= {"AguardandoLiberacao"}:
        return "auditoria"
    if status & {"Pendente", "Concluído", "Concluido"}:
        return "lancamento"
    return ""


def _situacao(itens, recusada: bool) -> str:
    if recusada:
        return "Recusada na auditoria"
    status = {str(i.status or "").strip() for i in itens}
    for chave, rotulo in (
        (_LANCADO, "Lançada"), ({"Concluído", "Concluido"}, "Aguardando lançamento"),
        ({"Devolvido"}, "Devolvida"), ({"AguardandoDevolucao"}, "Aguardando devolução"),
        ({"AguardandoLiberacao"}, "Em auditoria"),
    ):
        if status & chave:
            return rotulo
    return "Aguardando conferência"


def _ocs_da_nota(itens) -> list[str]:
    ocs = []
    for item in itens:
        for parte in re.split(r"[;,\n]+", str(item.pedido_compra or "")):
            parte = parte.strip()
            if parte and parte not in ocs:
                ocs.append(parte)
    return ocs


def _ref(item) -> str:
    """Identidade da nota, igual à da tela (_documento_ref_item)."""
    chave = re.sub(r"\D", "", str(item.chave_acesso or ""))
    if chave:
        return chave
    externo = str(item.documento_externo_id or "").strip()
    return f"doc:{externo}" if externo else f"item:{item.id}"


# --------------------------------------------------------------------------- planilha

def _aba_resumo(ws, f, dados, usuario):
    ws.title = "Resumo"
    notas, ocs, grv_ok = dados["notas"], dados["ocs"], dados["grv_ok"]
    entraram = [n for n in notas if n["entrou_no_periodo"]]
    lancadas = [n for n in notas if n["lancada_no_periodo"]]
    validas = [o for o in ocs if not o["cancelada"]]

    ws.append(["Relatório do Documento de Entrada"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Período: {f['inicio']:%d/%m/%Y} a {f['fim']:%d/%m/%Y}"])
    ws.append([f"Filtros: {_descrever_filtros(f)}"])
    ws.append([f"Gerado em {agora_br():%d/%m/%Y %H:%M} por {usuario}"])
    if not grv_ok:
        ws.append(["Atenção: o GRV não respondeu. As ordens de compra e o comprador ficaram de fora."])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C00000")
    elif not dados["itens_ok"]:
        ws.append(["Atenção: o GRV não devolveu os itens das OCs. Tente gerar o relatório de novo."])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C00000")
    ws.append([])

    _secao(ws, "Notas", ["", "Total"])
    _numeros(ws, [
        ("Notas que entraram no Sync", len(entraram), None),
        *((f"   {rotulo}", sum(1 for n in entraram if n["tipo"] == chave), None) for chave, rotulo in TIPOS.items()),
        ("Valor das notas que entraram", sum(n["valor"] for n in entraram), _FMT_MOEDA),
        ("   sem OC vinculada", sum(1 for n in entraram if not n["ocs"]), None),
        ("Auditadas e aprovadas", sum(1 for n in notas if n["auditada_no_periodo"] and n["auditoria"] == "Aprovada"), None),
        ("Recusadas na auditoria", sum(1 for n in notas if n["auditada_no_periodo"] and n["auditoria"] == "Recusada"), None),
        ("Conferidas fisicamente", sum(1 for n in notas if n["conferida_no_periodo"]), None),
        ("   com divergência na conferência", sum(1 for n in notas if n["conferida_no_periodo"] and n["divergencia"]), None),
        ("Lançadas no ERP", len(lancadas), None),
        ("Valor lançado", sum(n["valor"] for n in lancadas), _FMT_MOEDA),
        ("Estornos de lançamento", sum(len(n["estornos_periodo"]) for n in notas), None),
        ("Das que entraram, hoje em auditoria", sum(1 for n in entraram if n["etapa"] == "auditoria"), None),
        ("Das que entraram, hoje em lançamento", sum(1 for n in entraram if n["etapa"] == "lancamento"), None),
        ("Das que entraram, hoje finalizadas", sum(1 for n in entraram if n["etapa"] == "finalizado"), None),
        ("Dias da emissão até a entrada no Sync (média)", _media_dias((n["emissao"], n["entrou_em"]) for n in entraram), _FMT_DIAS),
        ("Dias da entrada até o lançamento (média)", _media_dias((n["entrou_em"], n["lancada_em"]) for n in lancadas), _FMT_DIAS),
    ])

    _secao(ws, "Compras (GRV)", ["", "Total"])
    if grv_ok:
        _numeros(ws, [
            ("OCs geradas", len(ocs), None),
            ("   canceladas", len(ocs) - len(validas), None),
            ("Valor das OCs geradas (sem as canceladas)", sum(o["total"] for o in validas), _FMT_MOEDA),
            ("   já recebidas", sum(1 for o in validas if o["situacao"] == "Recebida"), None),
            ("   recebidas em parte", sum(1 for o in validas if o["situacao"] == "Recebida em parte"), None),
            ("   em aberto", sum(1 for o in validas if o["situacao"] == "Em aberto"), None),
            ("Compradores com OC no período", len({o["comprador"] for o in ocs if o["comprador"]}), None),
            ("Dias da OC até a nota entrar no Sync (média)", _media_dias((n["data_oc"], n["entrou_em"]) for n in entraram), _FMT_DIAS),
        ])
    else:
        ws.append(["GRV indisponível"])
        ws.append([])

    if grv_ok:
        _tabela(ws, "Por comprador", ["", "OCs geradas", "Canceladas", "Valor das OCs", "Notas que entraram", "Notas lançadas"],
                _quebra_comprador(notas, ocs), {4: _FMT_MOEDA})
    if grv_ok and dados["itens_ok"]:
        _tabela(ws, "Custo dos itens por família e grupo", ["Família", "Grupo", "Itens", "Valor"],
                _quebra_familia(dados["itens"]), {4: _FMT_MOEDA})
        _tabela(ws, "Custo dos itens por quem solicitou", ["", "Itens", "OCs", "Valor"],
                _quebra_solicitante(dados["itens"]), {4: _FMT_MOEDA})
    _tabela(ws, "Por fornecedor", ["", "Notas que entraram", "Valor das notas", "Notas lançadas", "OCs geradas", "Valor das OCs"],
            _quebra_fornecedor(notas, ocs), {3: _FMT_MOEDA, 6: _FMT_MOEDA})
    _tabela(ws, "Por pessoa (quem fez)", ["", "Importou", "Auditou", "Conferiu", "Lançou"], _quebra_pessoa(notas))
    _tabela(ws, "Por tipo de documento", ["", "Entraram", "Valor que entrou", "Lançadas"],
            _quebra_tipo(notas), {3: _FMT_MOEDA})

    ws.column_dimensions["A"].width = 50
    for col in range(2, 7):
        ws.column_dimensions[get_column_letter(col)].width = 18


def _numeros(ws, numeros):
    for rotulo, valor, formato in numeros:
        _append(ws, [rotulo, valor if valor is not None else "Sem dado"])
        if formato and valor is not None:
            ws.cell(row=ws.max_row, column=2).number_format = formato
    ws.append([])


def _quebra_comprador(notas, ocs) -> list[list]:
    grupos = defaultdict(lambda: defaultdict(float))
    for o in ocs:
        g = grupos[o["comprador"] or _SEM_INFO]
        g["ocs"] += 1
        if o["cancelada"]:
            g["canceladas"] += 1
        else:
            g["valor"] += o["total"]
    for n in notas:
        chaves = n["compradores"] or ([_SEM_INFO] if n["ocs"] else [_SEM_OC])
        for chave in chaves:
            grupos[chave]["entraram"] += n["entrou_no_periodo"]
            grupos[chave]["lancadas"] += n["lancada_no_periodo"]
    linhas = [[k, int(g["ocs"]), int(g["canceladas"]), g["valor"], int(g["entraram"]), int(g["lancadas"])]
              for k, g in grupos.items() if g["ocs"] or g["entraram"] or g["lancadas"]]
    return sorted(linhas, key=lambda r: (-r[1], -r[4], r[0]))


def _quebra_familia(itens) -> list[list]:
    grupos = defaultdict(lambda: [0, 0.0])
    for i in itens:
        if i["conta"]:
            g = grupos[(i["familia"] or _SEM_INFO, i["grupo"] or _SEM_INFO)]
            g[0] += 1
            g[1] += i["total"]
    linhas = [[familia, grupo, g[0], g[1]] for (familia, grupo), g in grupos.items()]
    return sorted(linhas, key=lambda r: (-r[3], r[0], r[1]))


def _quebra_solicitante(itens) -> list[list]:
    # Item de várias solicitações sai com os nomes juntos ("A, B"): dividir
    # contaria o mesmo custo duas vezes.
    grupos = defaultdict(lambda: {"itens": 0, "ocs": set(), "valor": 0.0})
    for i in itens:
        if i["conta"]:
            g = grupos[i["solicitante"] or (_SEM_INFO if i["solicitacao"] else _SEM_SOLICITACAO)]
            g["itens"] += 1
            g["ocs"].add(i["oc"])
            g["valor"] += i["total"]
    linhas = [[k, g["itens"], len(g["ocs"]), g["valor"]] for k, g in grupos.items()]
    return sorted(linhas, key=lambda r: (-r[3], r[0]))


def _quebra_fornecedor(notas, ocs) -> list[list]:
    # Pelo CNPJ: o nome do XML e o do cadastro do GRV costumam vir diferentes.
    # Documento não fiscal vem sem CNPJ: cai no CNPJ de outra nota de mesmo nome.
    grupos = defaultdict(lambda: defaultdict(float))
    nomes = {}
    cnpj_por_nome = {}
    for registro in [*notas, *ocs]:
        if registro["cnpj"] and registro["fornecedor"]:
            cnpj_por_nome.setdefault(registro["fornecedor"].casefold(), registro["cnpj"])

    def chave(cnpj, nome):
        k = cnpj or cnpj_por_nome.get((nome or "").casefold()) or (nome or _SEM_INFO).casefold()
        nomes.setdefault(k, nome or _SEM_INFO)
        return k

    for n in notas:
        g = grupos[chave(n["cnpj"], n["fornecedor"])]
        if n["entrou_no_periodo"]:
            g["entraram"] += 1
            g["valor_notas"] += n["valor"]
        g["lancadas"] += n["lancada_no_periodo"]
    for o in ocs:
        g = grupos[chave(o["cnpj"], o["fornecedor"])]
        g["ocs"] += 1
        if not o["cancelada"]:
            g["valor_ocs"] += o["total"]
    linhas = [[nomes[k], int(g["entraram"]), g["valor_notas"], int(g["lancadas"]), int(g["ocs"]), g["valor_ocs"]]
              for k, g in grupos.items() if g["entraram"] or g["lancadas"] or g["ocs"]]
    return sorted(linhas, key=lambda r: (-(r[2] + r[5]), r[0]))


def _quebra_pessoa(notas) -> list[list]:
    grupos = defaultdict(lambda: defaultdict(int))
    for n in notas:
        if n["entrou_no_periodo"]:
            grupos[n["importada_por"] or _SEM_INFO]["importou"] += 1
        if n["auditada_no_periodo"]:
            grupos[n["auditada_por"] or _SEM_INFO]["auditou"] += 1
        if n["conferida_no_periodo"]:
            grupos[n["conferida_por"] or _SEM_INFO]["conferiu"] += 1
        if n["lancada_no_periodo"]:
            grupos[n["lancada_por"] or _SEM_INFO]["lancou"] += 1
    linhas = [[k, g["importou"], g["auditou"], g["conferiu"], g["lancou"]] for k, g in grupos.items()]
    return sorted(linhas, key=lambda r: (-sum(r[1:]), r[0]))


def _quebra_tipo(notas) -> list[list]:
    grupos = defaultdict(lambda: defaultdict(float))
    for n in notas:
        g = grupos[n["tipo_doc"]]
        if n["entrou_no_periodo"]:
            g["entraram"] += 1
            g["valor"] += n["valor"]
        g["lancadas"] += n["lancada_no_periodo"]
    linhas = [[k, int(g["entraram"]), g["valor"], int(g["lancadas"])] for k, g in grupos.items() if g["entraram"] or g["lancadas"]]
    return sorted(linhas, key=lambda r: (-r[1], r[0]))


def _aba_notas(ws, notas):
    colunas = [
        ("Tipo", "tipo_doc", None), ("Número", "numero", None), ("Fornecedor", "fornecedor", None),
        ("CNPJ", "cnpj", None), ("Emissão", "emissao", _FMT_DATA), ("Entrou no Sync em", "entrou_em", _FMT_DATA),
        ("Importada por", "importada_por", None), ("Situação atual", "situacao", None), ("Etapa atual", "etapa_nome", None),
        ("Auditoria", "auditoria", None), ("Auditada em", "auditada_em", _FMT_DATA), ("Auditada por", "auditada_por", None),
        ("Conferida em", "conferida_em", _FMT_DATA), ("Conferida por", "conferida_por", None),
        ("Sem conferência física", "sem_conf", None), ("Tentativas de conferência", "tentativas", None),
        ("Divergência na conferência", "div", None), ("Lançada em", "lancada_em", _FMT_DATA),
        ("Lançada por", "lancada_por", None), ("Nº lançamento", "numero_lancamento", None),
        ("Estornos de lançamento", "qtd_estornos", None), ("Valor", "valor", _FMT_MOEDA), ("Itens", "qtd_itens", None),
        ("OC", "oc", None), ("Comprador", "comprador", None), ("Data da OC", "data_oc", _FMT_DATA), ("CFOP", "cfops", None),
        ("Material de cliente", "mat_cliente", None), ("Remessa", "rem", None), ("Mercado Livre", "ml", None),
        ("Retorno de industrialização", "ret_ind", None), ("Manifestação", "manifestacao", None),
        ("Chave de acesso", "chave", None),
    ]

    def sim(valor):
        return "Sim" if valor else ""

    _lista(ws, colunas, [
        {**n, "etapa_nome": ETAPAS.get(n["etapa"], ""), "sem_conf": sim(n["sem_conferencia"]), "div": sim(n["divergencia"]),
         "qtd_estornos": len(n["estornos"]), "qtd_itens": len(n["itens"]), "oc": ", ".join(n["ocs"]),
         "mat_cliente": sim(n["material_cliente"]), "rem": sim(n["remessa"]), "ml": sim(n["mercado_livre"]),
         "ret_ind": sim(n["retorno_industrializacao"])}
        for n in notas
    ])


def _aba_ocs(ws, dados):
    if not dados["grv_ok"]:
        ws.append(["O GRV não respondeu. Tente gerar o relatório de novo em instantes."])
        ws.column_dimensions["A"].width = 70
        return
    colunas = [
        ("OC", "codigo", None), ("Data da OC", "data_oc", _FMT_DATA), ("Fornecedor", "fornecedor", None), ("CNPJ", "cnpj", None),
        ("Comprador", "comprador", None), ("Status no GRV", "status_oc", None), ("Situação", "situacao", None),
        ("Valor", "total", _FMT_MOEDA), ("Itens", "qtd_itens", None), ("Itens com saldo a receber", "itens_com_saldo", None),
        ("Condição de pagamento", "cond_pagamento", None), ("Forma de pagamento", "forma_pgto", None),
        ("Notas no Sync", "notas", None),
    ]
    _lista(ws, colunas, dados["ocs"])


def _aba_itens_oc(ws, dados):
    if not dados["grv_ok"] or not dados["itens_ok"]:
        ws.append(["O GRV não respondeu. Tente gerar o relatório de novo em instantes."])
        ws.column_dimensions["A"].width = 70
        return
    colunas = [
        ("OC", "oc", None), ("Data da OC", "data_oc", _FMT_DATA), ("Fornecedor", "fornecedor", None),
        ("Comprador", "comprador", None), ("Tipo", "tipo", None), ("Item", "item", None), ("Código", "codigo", None),
        ("Descrição", "descricao", None), ("Família", "familia", None), ("Grupo", "grupo", None),
        ("Unidade", "unidade", None), ("Qtd", "qtde", None), ("Valor unitário", "valor_unitario", _FMT_UNITARIO),
        ("Total", "total", _FMT_MOEDA), ("IPI", "ipi", _FMT_MOEDA), ("Solicitado por", "solicitante", None),
        ("Solicitação", "solicitacao", None), ("Solicitada em", "solicitada_em", _FMT_DIA), ("OS", "os", None),
        ("Situação", "situacao", None),
    ]
    _lista(ws, colunas, [
        {**i, "solicitante": i["solicitante"] or ("" if i["solicitacao"] else _SEM_SOLICITACAO)}
        for i in dados["itens"]
    ])


# --------------------------------------------------------------------------- apoio

def _media_dias(pares):
    dias = [(fim - ini).total_seconds() / 86400 for ini, fim in pares if ini and fim and fim >= ini]
    return round(sum(dias) / len(dias), 1) if dias else None


def _contem(texto, busca) -> bool:
    return busca.casefold() in str(texto or "").casefold()


def _num(texto):
    texto = str(texto or "")
    return (0, int(texto), "") if texto.isdigit() else (1, 0, texto)


def _float(valor) -> float:
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def _dt(valor):
    """A bridge devolve data em texto ISO; o Postgres direto, em date/datetime."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    if isinstance(valor, date):
        return datetime.combine(valor, time.min)
    try:
        return datetime.fromisoformat(str(valor).strip().replace("Z", "")).replace(tzinfo=None)
    except ValueError:
        return None


def _descrever_filtros(f) -> str:
    partes = []
    if f["tipo"]:
        partes.append(f"tipo {TIPOS[f['tipo']]}")
    if f["etapa"]:
        partes.append(f"etapa atual {ETAPAS[f['etapa']]}")
    if f["fornecedor"]:
        partes.append(f"fornecedor contém \"{f['fornecedor']}\"")
    if f["comprador"]:
        partes.append(f"comprador contém \"{f['comprador']}\"")
    texto = ", ".join(partes) or "nenhum (tudo)"
    if f["tipo"] or f["etapa"]:
        texto += " - tipo e etapa valem só para as notas"
    return texto
