"""Cardex: itens/famílias desconsiderados (aparecem, mas não entram nos
totais) e o arredondamento em centavos por lançamento. Mesma fixture do
test_logistica_cardex; bridge sempre com patch."""
from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from conferencia_app.extensions import db
from conferencia_app.models import LogisticaCardexDesconsiderado, PermissaoAcesso
from conferencia_app.services import logistica_cardex_service as svc
from tests.test_app import build_test_app, set_logged_user
from tests.test_logistica_cardex import (
    MOV_ENTRADA_78851, MOV_ENTRADA_173975, MOV_SAIDA_53918, NOTAS, VALOR_HOJE,
    _app_logado, _grv, _mov, _produto,
)

SET = "inicio=2026-09-01&fim=2026-09-30"
API = "/api/logistica/cardex"
COD_B = 5002


@pytest.fixture(autouse=True)
def _sem_cache():
    svc.limpar_cache()
    yield
    svc.limpar_cache()


def _grv_duas_familias():
    """A chapa da fixture (matéria-prima) + um item de outra família, parado."""
    outro = {"cod_produto": COD_B, "codigo_interno": "30-02-00010", "codigo": "300200010",
             "descricao": "PARAFUSO M8", "unidade": "PC", "familia": "N - 02 - CONSUMO", "grupo": "2",
             "preco_custo": 2.5, "fiscal": {"classificacao_fiscal": "73181500"}, "saldos": {"1": 100.0}}
    return _grv(produtos=[_produto(), outro])


def _resumo(client, extra=""):
    return client.get(f"{API}/resumo?{SET}{extra}").get_json()


def test_desconsiderar_item_tira_do_total_mas_continua_na_lista(tmp_path):
    app, client = _app_logado(tmp_path)
    with patch.object(svc, "buscar_cardex_grv", return_value=_grv_duas_familias()):
        antes = _resumo(client)
        assert antes["indicadores"]["itens"] == 2 and antes["indicadores"]["itens_desconsiderados"] == 0
        assert antes["totais"]["valor_final"] == pytest.approx(VALOR_HOJE + 250.0, abs=0.01)

        resp = client.post(f"{API}/desconsiderados", json={
            "tipo": "ITEM", "chave": "30-02-00010", "motivo": "Material de consumo, não é estoque", "descricao": "PARAFUSO M8"})
        assert resp.status_code == 201, resp.get_json()
        registro = resp.get_json()["desconsiderado"]
        assert registro["chave"] == "300200010" and registro["por"] == "ADMIN"

        depois = _resumo(client)
        # Continua na lista, com a marca; sai de total, subtotal e indicador.
        assert len(depois["linhas"]) == 2
        marcado = next(l for l in depois["linhas"] if l["codigo"] == "300200010")
        assert marcado["desconsiderado"]["motivo"] == "Material de consumo, não é estoque"
        assert marcado["valor_final"] == pytest.approx(250.0)
        assert depois["totais"]["valor_final"] == pytest.approx(VALOR_HOJE, abs=0.01)
        assert depois["indicadores"]["itens"] == 1
        assert depois["indicadores"]["itens_desconsiderados"] == 1
        assert depois["indicadores"]["valor_desconsiderado"] == pytest.approx(250.0)
        assert depois["desconsiderados"]["totais"]["valor_final"] == pytest.approx(250.0)
        assert depois["subtotais_familia"]["N - 02 - CONSUMO"]["valor_final"] == 0

        item = client.get(f"{API}/item?{SET}&codigo=30-02-00010").get_json()
        assert item["desconsiderado"]["id"] == registro["id"]
        assert client.get(f"{API}/desconsiderados").get_json()["desconsiderados"][0]["chave"] == "300200010"

        # Desfazer: volta para o total e fica o histórico de quem desfez.
        assert client.delete(f"{API}/desconsiderados/{registro['id']}").status_code == 200
        assert _resumo(client)["totais"]["valor_final"] == pytest.approx(VALOR_HOJE + 250.0, abs=0.01)
        assert client.get(f"{API}/desconsiderados").get_json()["desconsiderados"] == []
    with app.app_context():
        reg = LogisticaCardexDesconsiderado.query.one()
        assert reg.ativo is False and reg.removido_por == "ADMIN" and reg.removido_em is not None


def test_desconsiderar_familia_inteira(tmp_path):
    _, client = _app_logado(tmp_path)
    with patch.object(svc, "buscar_cardex_grv", return_value=_grv_duas_familias()):
        resp = client.post(f"{API}/desconsiderados", json={
            "tipo": "FAMILIA", "chave": "N - 01 - MATÉRIA-PRIMA", "motivo": "Família controlada fora do Cardex"})
        assert resp.status_code == 201, resp.get_json()
        dados = _resumo(client)
        assert dados["totais"]["valor_final"] == pytest.approx(250.0)
        assert dados["indicadores"]["itens_desconsiderados"] == 1
        fam = {f["familia"]: f for f in dados["familias_resumo"]}
        assert fam["N - 01 - MATÉRIA-PRIMA"]["desconsiderada"]["tipo"] == "FAMILIA"
        assert fam["N - 01 - MATÉRIA-PRIMA"]["valor_desconsiderado"] == pytest.approx(VALOR_HOJE, abs=0.01)
        assert fam["N - 01 - MATÉRIA-PRIMA"]["totais"]["valor_final"] == 0
        assert fam["N - 02 - CONSUMO"]["desconsiderada"] is None
        # O painel por família não some quando a tabela é filtrada por família.
        filtrado = _resumo(client, "&familia=N - 02 - CONSUMO")
        assert len(filtrado["linhas"]) == 1 and len(filtrado["familias_resumo"]) == 2


def test_desconsiderar_valida_entrada_e_permissao(tmp_path):
    app, client = _app_logado(tmp_path)
    ok = {"tipo": "ITEM", "chave": "19-01-00564", "motivo": "Motivo válido"}
    assert client.post(f"{API}/desconsiderados", json={**ok, "motivo": "abc"}).status_code == 400
    assert client.post(f"{API}/desconsiderados", json={**ok, "tipo": "GRUPO"}).status_code == 400
    assert client.post(f"{API}/desconsiderados", json={**ok, "chave": "  "}).status_code == 400
    assert client.post(f"{API}/desconsiderados", json=ok).status_code == 201
    assert client.post(f"{API}/desconsiderados", json=ok).status_code == 400  # já desconsiderado
    assert client.delete(f"{API}/desconsiderados/999999").status_code == 404

    # Sem acesso à página: não desconsidera.
    set_logged_user(client, "LOGISTICA_TESTE", "Logística")
    outro = {"tipo": "ITEM", "chave": "30-02-00010", "motivo": "Motivo válido"}
    assert client.post(f"{API}/desconsiderados", json=outro).status_code == 403

    # Com acesso só à página (sem a permissão de fechar mês): desconsidera.
    with app.app_context():
        db.session.add(PermissaoAcesso(scope_type="USER", scope_id="LOGISTICA_TESTE",
                                       permission_key="PAGE_LOGISTICA_CARDEX", allow=True))
        db.session.commit()
    assert client.post(f"{API}/desconsiderados", json=outro).status_code == 201
    assert client.post(f"{API}/fechamentos", json={"ano": 2026, "mes": 8}).status_code == 403


def test_mes_fechado_guarda_a_marca_e_nao_muda_depois(tmp_path):
    _, client = _app_logado(tmp_path)
    with patch.object(svc, "buscar_cardex_grv", return_value=_grv_duas_familias()):
        registro = client.post(f"{API}/desconsiderados", json={
            "tipo": "ITEM", "chave": "30-02-00010", "motivo": "Fora do inventário contábil"}).get_json()["desconsiderado"]
        resp = client.post(f"{API}/fechamentos", json={"ano": 2026, "mes": 8})
        assert resp.status_code == 201, resp.get_json()
        fech = resp.get_json()["fechamento"]
    # O total do fechamento não inclui o desconsiderado, mas o item está na foto.
    assert fech["total_itens"] == 2
    assert fech["valor_final"] == pytest.approx(18707.09, abs=0.01)
    assert fech["desconsiderados"] == {"itens": 1, "valor_final": pytest.approx(250.0)}
    assert any("desconsiderado" in a for a in fech["avisos"])

    # Voltar a considerar depois do fechamento não mexe no mês fechado.
    assert client.delete(f"{API}/desconsiderados/{registro['id']}").status_code == 200
    with patch.object(svc, "buscar_cardex_grv", side_effect=AssertionError("não devia ir ao GRV")):
        agosto = client.get(f"{API}/resumo?inicio=2026-08-01&fim=2026-08-31").get_json()
    assert agosto["fonte"] == "congelado"
    assert agosto["indicadores"]["itens_desconsiderados"] == 1
    assert agosto["totais"]["valor_final"] == pytest.approx(18707.09, abs=0.01)
    # ...mas o mês seguinte, ao vivo, já considera o item de novo.
    with patch.object(svc, "buscar_cardex_grv", return_value=_grv_duas_familias()):
        setembro = _resumo(client, "&atualizar=1")
    assert setembro["indicadores"]["itens_desconsiderados"] == 0
    assert setembro["totais"]["valor_final"] == pytest.approx(VALOR_HOJE + 250.0, abs=0.01)


def test_exportacao_separa_os_desconsiderados(tmp_path):
    _, client = _app_logado(tmp_path)
    with patch.object(svc, "buscar_cardex_grv", return_value=_grv_duas_familias()):
        client.post(f"{API}/desconsiderados", json={"tipo": "ITEM", "chave": "30-02-00010", "motivo": "Fora do inventário"})
        resp = client.get(f"{API}/exportar.xlsx?{SET}&tipo=resumo")
        assert resp.status_code == 200
        linhas = list(load_workbook(BytesIO(resp.data))["Resumo"].iter_rows(values_only=True))
        rotulos = [l[0] for l in linhas]
        total = rotulos.index("TOTAL GERAL")
        secao = next(i for i, r in enumerate(rotulos) if str(r or "").startswith("ITENS DESCONSIDERADOS"))
        assert secao > total
        assert linhas[total][17] == pytest.approx(VALOR_HOJE, abs=0.01)  # Saldo final R$ sem o desconsiderado
        fora = next(l for l in linhas[secao:] if l[1] == "30-02-00010")
        assert fora[17] == pytest.approx(250.0) and "Fora do inventário" in fora[19]
        assert any(str(l[0] or "").startswith("TOTAL DESCONSIDERADO") for l in linhas[secao:])
        pdf = client.get(f"{API}/exportar.pdf?{SET}&tipo=resumo")
        assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")


# --------------------------------------------------------------------------
# Arredondamento em centavos por lançamento
# --------------------------------------------------------------------------

def test_saldo_inicial_mais_movimentos_fecha_exato_no_saldo_final():
    grv = _grv()
    item = svc.calcular_item(grv["produtos"][0], grv["movimentos"], grv["notas"],
                             inicio=svc.date(2026, 9, 1), fim=svc.date(2026, 9, 30))
    det, res = svc.detalhe(item), svc.resumo(item)
    for linha in det["linhas"]:
        assert linha["valor"] == round(linha["valor"], 2)
        assert linha["saldo_valor"] == round(linha["saldo_valor"], 2)
    soma = round(det["inicial"]["valor"] + sum(l["valor"] for l in det["linhas"]), 2)
    assert soma == det["final"]["valor"]
    assert res["valor_reavaliacao"] == 0
    assert round(res["valor_inicial"] + res["valor_entrada"] + res["valor_saida"], 2) == res["valor_final"]


def test_saida_que_zera_o_estoque_leva_o_centavo_que_sobrou():
    # 3 un. valendo R$ 10,00 (médio 3,3333...): três saídas de 1 un. dariam
    # 3,33 + 3,33 + 3,33 = 9,99 e sobraria 0,01 pendurado num saldo zero.
    produto = _produto(saldos={"1": 0.0}, preco=0.0)
    saidas = [_mov(10 + i, f"2026-09-0{i + 1}", 1, -1.0, "TSAIDA_E", f"1;{900 + i}") for i in range(3)]
    item = svc.calcular_item(produto, saidas, {}, inicio=svc.date(2026, 9, 1), fim=svc.date(2026, 9, 30),
                             abertura_ancora={"q": {"1": 3.0}, "m": 10.0 / 3})
    det = svc.detalhe(item)
    assert det["inicial"]["valor"] == 10.0
    assert [l["valor"] for l in det["linhas"]] == [-3.33, -3.33, -3.34]
    # O médio é sempre saldo R$ / saldo qtde: depois de duas saídas sobra
    # 1 un. valendo 3,34, e é por esse médio que a última sai.
    assert det["linhas"][-1]["custo_unit"] == pytest.approx(3.34)
    assert det["final"]["qtde"] == 0.0 and det["final"]["valor"] == 0.0
    assert svc.resumo(item)["valor_reavaliacao"] == 0
